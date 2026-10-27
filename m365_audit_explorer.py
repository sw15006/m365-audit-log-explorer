#!/usr/bin/env python3
"""
Microsoft 365 Audit Log Explorer  v1.1
--------------------------------------
Fetch, search, summarise and download the Microsoft 365 unified audit log
(Entra ID, Exchange, SharePoint / OneDrive, Teams, Power BI, Compliance, DLP)
through the Office 365 Management Activity API.

v1.1 adds:
  - Detections tab (password spray, brute force, mass download / deletion)
  - MITRE ATT&CK technique on every quick hunt and detection
  - New hunts (Conditional Access, federation, guests, Power Automate, ...)
  - IP Intel tab (country / ASN via ip-api.com, optional AbuseIPDB check)
  - User Timeline tab (one user's activity in time order, with hunt flags)
  - Warnings when content blobs fail or the window is clipped to 7 days

Developed by Siddhesh Nagarkar - SOC Team

1. pip install requests            (and: pip install openpyxl  for Excel export)
2. python m365_audit_explorer.py
3. Enter the connection details in the Connect window.

Connection details are entered at every start and kept in memory only.
The client secret is never written to disk - not in profiles, logs or output.

Azure requirement: the app registration needs the API permission
"Office 365 Management APIs" -> Application -> ActivityFeed.Read
(plus ActivityFeed.ReadDlp for DLP events), with admin consent.
"""

# ==========================================================================
#            OPTIONAL SETTINGS (no credentials are stored here)
# ==========================================================================

DEFAULT_BASE_URL  = "https://manage.office.com/api/v1.0"
DEFAULT_TOKEN_URL = "https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token"

OUTPUT_DIR       = "output_m365_audit"   # default download folder
DEFAULT_FORMAT   = "csv"                 # csv / json / jsonl / xlsx
VERIFY_SSL       = True                  # False if an SSL-inspection proxy causes errors
PROXY            = ""                    # e.g. "http://proxy.company.local:8080"
DOWNLOAD_THREADS = 6                     # parallel content downloads
MAX_EVENTS_IN_MEMORY = 300000            # "Fetch events" limit; use "Download directly" for more
LATE_ARRIVAL_HOURS   = 12                # look this far past the window end for late events
RATE_PER_MINUTE  = 1500                  # API allows 2000/min per tenant

# Detections tab defaults (can also be changed in the tab itself)
SPRAY_MIN_USERS     = 10    # one IP fails sign-in for at least this many different users ...
SPRAY_WINDOW_MIN    = 60    # ... within this many minutes
BRUTE_MIN_FAILS     = 5     # failed sign-ins for one user ...
BRUTE_WINDOW_MIN    = 30    # ... within this many minutes
MASS_DOWNLOAD_FILES = 100   # files downloaded by one user ...
MASS_DELETE_FILES   = 100   # files / folders deleted by one user ...
MASS_WINDOW_MIN     = 10    # ... within this many minutes

# IP Intel tab - lookups only run when you click the buttons; nothing is sent automatically
GEO_LOOKUP_URL    = "http://ip-api.com/batch"   # free, no key, HTTP only (Pro plan needed for HTTPS)
ABUSEIPDB_URL     = "https://api.abuseipdb.com/api/v2/check"
ABUSEIPDB_MAX_IPS = 200     # max IPs checked per click (free plan: 1000 checks/day)

PROFILES_FILE     = ""   # blank = m365_profiles.json next to the script
ACTIVITY_LOG_FILE = ""   # blank = m365_activity_log.csv next to the script

# ==========================================================================
#                     NO NEED TO EDIT BELOW THIS LINE
# ==========================================================================

import csv
import getpass
import ipaddress
import json
import os
import platform
import queue
import re
import subprocess
import sys
import threading
import time
import zipfile
from collections import Counter, deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta, timezone
from urllib.parse import urlparse

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog

import requests

TOOL_NAME = "Microsoft 365 Audit Log Explorer"
TOOL_VERSION = "1.1"
DEVELOPER = "Siddhesh Nagarkar"
TEAM = "SOC Team"

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROFILES_PATH = PROFILES_FILE or os.path.join(SCRIPT_DIR, "m365_profiles.json")
ACTIVITY_PATH = ACTIVITY_LOG_FILE or os.path.join(SCRIPT_DIR, "m365_activity_log.csv")
PROXIES = {"http": PROXY, "https": PROXY} if PROXY else None

MAX_GRID_ROWS = 20000
IST = timedelta(hours=5, minutes=30)
ISO_RE = re.compile(r"^(\d{4}-\d{2}-\d{2})T(\d{2}:\d{2}:\d{2})")
RETENTION = timedelta(days=7)

CONTENT_TYPES = [
    ("Audit.AzureActiveDirectory", "Entra ID (Azure AD)", True),
    ("Audit.Exchange", "Exchange", True),
    ("Audit.SharePoint", "SharePoint / OneDrive", True),
    ("Audit.General", "General (Teams, Power BI, Compliance...)", True),
    ("DLP.All", "DLP (needs ActivityFeed.ReadDlp)", False),
]

PRESETS = [
    ("Last 15 minutes", 15), ("Last 20 minutes", 20), ("Last 30 minutes", 30),
    ("Last 1 hour", 60), ("Last 4 hours", 240), ("Last 12 hours", 720),
    ("Last 24 hours", 1440), ("Last 3 days", 4320), ("Last 7 days", 10080),
    ("Custom minutes", "minutes"), ("Custom range (UTC)", "range"),
]

PREFERRED_COLS = ["CreationTime", "Workload", "Operation", "UserId", "ClientIP", "ClientIPAddress",
                  "ResultStatus", "RecordType", "ObjectId", "MailboxOwnerUPN", "SiteUrl",
                  "SourceFileName", "UserAgent", "Id"]
STREAM_COLS = ["CreationTime", "Workload", "RecordType", "Operation", "UserId", "ClientIP",
               "ResultStatus", "ObjectId", "Id", "AuditData"]
WIDE_COLS = {"CreationTime", "UserId", "Operation", "ObjectId", "SiteUrl", "Value"}

RECORD_TYPES = {
    1: "ExchangeAdmin", 2: "ExchangeItem", 3: "ExchangeItemGroup", 4: "SharePoint",
    6: "SharePointFileOperation", 7: "OneDrive", 8: "AzureActiveDirectory",
    9: "AzureActiveDirectoryAccountLogon", 14: "SharePointSharingOperation",
    15: "AzureActiveDirectoryStsLogon", 18: "SecurityComplianceCenterEOPCmdlet",
    25: "MicrosoftTeams", 28: "ThreatIntelligence", 50: "ExchangeItemAggregated",
}

GROUP_FIELDS = [("Operation", "Operation"), ("User", "UserId"), ("Client IP", "_ip"),
                ("Country", "_country"), ("Workload", "Workload"), ("Result", "ResultStatus"),
                ("Record type", "RecordType")]

# --- MITRE ATT&CK techniques used by hunts and detections -----------------
MITRE = {
    "T1020": "Automated Exfiltration",
    "T1070.008": "Clear Mailbox Data",
    "T1078.004": "Valid Accounts: Cloud Accounts",
    "T1098": "Account Manipulation",
    "T1098.001": "Additional Cloud Credentials",
    "T1098.002": "Additional Email Delegate Permissions",
    "T1098.003": "Additional Cloud Roles",
    "T1110": "Brute Force",
    "T1110.001": "Password Guessing",
    "T1110.003": "Password Spraying",
    "T1114.002": "Remote Email Collection",
    "T1114.003": "Email Forwarding Rule",
    "T1136.003": "Create Account: Cloud Account",
    "T1213": "Data from Information Repositories",
    "T1213.002": "Data from Information Repositories: SharePoint",
    "T1484.002": "Domain or Tenant Policy Modification: Trust Modification",
    "T1485": "Data Destruction",
    "T1528": "Steal Application Access Token",
    "T1530": "Data from Cloud Storage",
    "T1556.006": "Modify Authentication Process: Multi-Factor Authentication",
    "T1562.001": "Impair Defenses: Disable or Modify Tools",
    "T1562.008": "Impair Defenses: Disable or Modify Cloud Logs",
    "T1564.008": "Hide Artifacts: Email Hiding Rules",
    "T1567": "Exfiltration Over Web Service",
}


def mitre_text(ids, names=False):
    if not ids:
        return ""
    if names:
        return "\n".join(f"{t}  {MITRE.get(t, '')}" for t in ids)
    return ", ".join(ids)


# --- Quick hunts: ready-made filters for common SOC checks ----------------
# operations = any of (exact names); text = event JSON contains any of (AND with operations)
QUICK_HUNTS = [
    {"category": "Exchange / Mailbox", "name": "Inbox rules created or changed",
     "description": "New or modified inbox rules - attackers use them to hide or forward mail (BEC).",
     "mitre": ["T1564.008", "T1114.003"],
     "operations": ["New-InboxRule", "Set-InboxRule", "UpdateInboxRules", "Enable-InboxRule"]},
    {"category": "Exchange / Mailbox", "name": "Mail forwarding configured",
     "description": "Mailbox, inbox-rule or transport-rule changes that forward or redirect mail outside.",
     "mitre": ["T1114.003"],
     "operations": ["Set-Mailbox", "New-InboxRule", "Set-InboxRule", "UpdateInboxRules",
                    "New-TransportRule", "Set-TransportRule"],
     "text": ["ForwardingSmtpAddress", "ForwardingAddress", "ForwardTo", "RedirectTo",
              "ForwardAsAttachmentTo", "BlindCopyTo"]},
    {"category": "Exchange / Mailbox", "name": "Mailbox permissions granted",
     "description": "Full access, send-as or folder permissions given to another account.",
     "mitre": ["T1098.002"],
     "operations": ["Add-MailboxPermission", "Add-RecipientPermission", "Add-MailboxFolderPermission",
                    "Set-MailboxFolderPermission"]},
    {"category": "Exchange / Mailbox", "name": "Mail items accessed",
     "description": "Mailbox read access (needs Audit Premium / E5). Useful to scope a compromised mailbox.",
     "mitre": ["T1114.002"],
     "operations": ["MailItemsAccessed"]},
    {"category": "Exchange / Mailbox", "name": "Mail deleted",
     "description": "Messages moved to Deleted Items, soft-deleted or hard-deleted.",
     "mitre": ["T1070.008"],
     "operations": ["HardDelete", "SoftDelete", "MoveToDeletedItems"]},
    {"category": "Exchange / Mailbox", "name": "Mail flow (transport) rule changes",
     "description": "Organisation-wide mail flow rules created, changed, disabled or removed.",
     "mitre": ["T1114.003", "T1564.008"],
     "operations": ["New-TransportRule", "Set-TransportRule", "Remove-TransportRule",
                    "Disable-TransportRule", "Enable-TransportRule"]},
    {"category": "Exchange / Mailbox", "name": "Mailbox auditing disabled or bypassed",
     "description": "Mailbox or organisation audit settings turned off, or an account set to bypass auditing.",
     "mitre": ["T1562.008"],
     "operations": ["Set-Mailbox", "Set-OrganizationConfig", "Set-MailboxAuditBypassAssociation"],
     "text": ["AuditEnabled", "AuditDisabled", "AuditBypassEnabled"]},
    {"category": "Exchange / Mailbox", "name": "Anti-phish / spam / malware protection weakened",
     "description": "Email protection policies or rules changed, disabled or removed, and new allow-list "
                    "entries - attackers use these to let phishing through.",
     "mitre": ["T1562.001"],
     "operations": ["Set-AntiPhishPolicy", "Remove-AntiPhishPolicy", "Disable-AntiPhishRule",
                    "Set-HostedContentFilterPolicy", "Set-HostedContentFilterRule",
                    "Disable-HostedContentFilterRule", "Remove-HostedContentFilterRule",
                    "Set-MalwareFilterPolicy", "Disable-MalwareFilterRule", "Set-SafeLinksPolicy",
                    "Set-SafeAttachmentPolicy", "New-TenantAllowBlockListItems",
                    "Set-TenantAllowBlockListItems"]},
    {"category": "Identity (Entra ID)", "name": "Failed sign-ins",
     "description": "Failed user logins - look for many failures per user or per IP. "
                    "See the Detections tab for automatic spray / brute-force checks.",
     "mitre": ["T1110"],
     "operations": ["UserLoginFailed"]},
    {"category": "Identity (Entra ID)", "name": "Successful sign-ins",
     "description": "Successful user logins - group by Client IP or Country to spot unusual locations.",
     "mitre": ["T1078.004"],
     "operations": ["UserLoggedIn"]},
    {"category": "Identity (Entra ID)", "name": "Admin role changes",
     "description": "Users added to or removed from directory roles.",
     "mitre": ["T1098.003"],
     "operations": ["Add member to role.", "Remove member from role.", "Add eligible member to role."]},
    {"category": "Identity (Entra ID)", "name": "User account changes",
     "description": "Users created, deleted, disabled, or passwords reset / changed.",
     "mitre": ["T1136.003", "T1098"],
     "operations": ["Add user.", "Delete user.", "Disable account.", "Reset user password.",
                    "Change user password.", "Set force change user password."]},
    {"category": "Identity (Entra ID)", "name": "Guest / external users invited or added",
     "description": "External users invited, redeemed or created as guests - check who invited them.",
     "mitre": ["T1136.003"],
     "operations": ["Invite external user.", "Redeem external user invite.", "Add user."],
     "text": ["#EXT#", "Guest", "external"]},
    {"category": "Identity (Entra ID)", "name": "MFA / security info changes",
     "description": "Authentication methods registered, changed or deleted.",
     "mitre": ["T1556.006"],
     "operations": ["User registered security info", "User deleted security info",
                    "User changed default security info", "Admin registered security info",
                    "Admin deleted security info", "Update user."],
     "text": ["StrongAuthentication", "security info"]},
    {"category": "Identity (Entra ID)", "name": "Conditional Access policy changes",
     "description": "Conditional Access policies added, changed or deleted. 'Update policy.' can also "
                    "include other directory policies - open the record to check the policy type.",
     "mitre": ["T1556.006", "T1562.001"],
     "operations": ["Add conditional access policy", "Update conditional access policy",
                    "Delete conditional access policy", "Add policy.", "Update policy.", "Delete policy."]},
    {"category": "Identity (Entra ID)", "name": "Domain and federation changes",
     "description": "Domains added or verified, or federation / authentication settings changed - "
                    "a backdoor federation trust lets an attacker sign in as any user (Golden SAML style).",
     "mitre": ["T1484.002"],
     "operations": ["Set domain authentication.", "Set federation settings on domain.",
                    "Add domain to company.", "Add unverified domain.", "Verify domain.",
                    "Remove domain from company.", "Update domain."]},
    {"category": "Identity (Entra ID)", "name": "App consent and permissions",
     "description": "Consent grants, new service principals, credentials and app role assignments.",
     "mitre": ["T1528", "T1098.001"],
     "operations": ["Consent to application.", "Add delegated permission grant.",
                    "Add app role assignment to service principal.", "Add service principal.",
                    "Add service principal credentials.", "Update application - Certificates and secrets management "]},
    {"category": "SharePoint / OneDrive", "name": "External / anonymous sharing",
     "description": "Sharing links and invitations, including anonymous 'anyone' links.",
     "mitre": ["T1567"],
     "operations": ["SharingSet", "SharingInvitationCreated", "AnonymousLinkCreated", "AnonymousLinkUsed",
                    "AddedToSecureLink", "SecureLinkCreated", "CompanyLinkCreated"]},
    {"category": "SharePoint / OneDrive", "name": "File downloads",
     "description": "Files downloaded or synced - group by User to find mass downloads.",
     "mitre": ["T1530", "T1213.002"],
     "operations": ["FileDownloaded", "FileSyncDownloadedFull"]},
    {"category": "SharePoint / OneDrive", "name": "File and folder deletions",
     "description": "Files or folders deleted or sent to the recycle bin - group by User for mass deletion.",
     "mitre": ["T1485"],
     "operations": ["FileDeleted", "FileRecycled", "FileDeletedFirstStageRecycleBin",
                    "FileDeletedSecondStageRecycleBin", "FolderDeleted", "FolderRecycled"]},
    {"category": "SharePoint / OneDrive", "name": "Site collection admin changes",
     "description": "Site collection administrators added or removed - gives full control of a site.",
     "mitre": ["T1098"],
     "operations": ["SiteCollectionAdminAdded", "SiteCollectionAdminRemoved"]},
    {"category": "Teams", "name": "Team and membership changes",
     "description": "Teams created or deleted, members added or removed, role changes.",
     "mitre": ["T1098"],
     "operations": ["TeamCreated", "TeamDeleted", "MemberAdded", "MemberRemoved", "MemberRoleChanged"]},
    {"category": "Power Platform", "name": "Power Automate flows created or changed",
     "description": "Flows and connections created or edited - flows can silently copy mail or files "
                    "to outside services.",
     "mitre": ["T1020"],
     "operations": ["CreateFlow", "EditFlow", "PutConnection", "PutPermissions"]},
    {"category": "Admin & Compliance", "name": "eDiscovery / content searches",
     "description": "Content searches created, run, previewed or exported.",
     "mitre": ["T1213", "T1114.002"],
     "operations": ["SearchCreated", "SearchStarted", "SearchExported", "SearchExportDownloaded",
                    "SearchPreviewed", "CaseAdded"]},
    {"category": "Admin & Compliance", "name": "Audit configuration changes",
     "description": "Changes to admin audit logging or mailbox audit bypass - possible defence evasion.",
     "mitre": ["T1562.008"],
     "operations": ["Set-AdminAuditLogConfig", "Set-MailboxAuditBypassAssociation", "Set-OrganizationConfig"]},
    {"category": "Admin & Compliance", "name": "DLP rule matches",
     "description": "Data loss prevention policy matches (requires the DLP.All content type).",
     "operations": ["DlpRuleMatch", "DlpRuleUndo"]},
    {"category": "Admin & Compliance", "name": "Failed operations (any workload)",
     "description": "Any event whose result status is a failure.",
     "result": "fail"},
]


# ==========================================================================
#                              API CLIENT
# ==========================================================================
class M365Error(Exception):
    def __init__(self, msg, status=None, code=""):
        super().__init__(msg)
        self.status, self.code = status, code


class RateLimiter:
    def __init__(self, per_min):
        self.per_min = per_min
        self.calls = deque()
        self.lock = threading.Lock()

    def acquire(self, stop=None):
        while True:
            with self.lock:
                now = time.time()
                while self.calls and now - self.calls[0] > 60:
                    self.calls.popleft()
                if len(self.calls) < self.per_min:
                    self.calls.append(now)
                    return
                wait = 60 - (now - self.calls[0]) + 0.1
            for _ in range(max(1, int(min(wait, 2) * 10))):
                if stop is not None and stop.is_set():
                    raise M365Error("Cancelled")
                time.sleep(0.1)


class M365Client:
    """Office 365 Management Activity API client. Credentials held in memory only."""

    def __init__(self, creds):
        self.tenant_id = creds["tenant_id"]
        base = (creds.get("base_url") or DEFAULT_BASE_URL).rstrip("/")
        self.feed = f"{base}/{self.tenant_id}/activity/feed"
        token_url = creds.get("token_url") or DEFAULT_TOKEN_URL
        self.token_url = token_url.replace("{tenant_id}", self.tenant_id).replace("<tenant_id>", self.tenant_id)
        self.scope = f"https://{urlparse(base).netloc}/.default"
        self._client_id = creds["client_id"]
        self._client_secret = creds["client_secret"]
        self._token = None
        self._expires = 0
        self._tlock = threading.Lock()
        self.limiter = RateLimiter(RATE_PER_MINUTE)
        self.log = None
        self.stop = None
        self.on_call = None      # callback(action, detail, count, error) for the activity log

    def wipe(self):
        self._client_secret = None
        self._token = None

    def token(self):
        with self._tlock:
            if self._token and time.time() < self._expires - 300:
                return self._token
            if not self._client_secret:
                raise M365Error("Not connected (credentials cleared)")
            data = {"grant_type": "client_credentials", "client_id": self._client_id,
                    "client_secret": self._client_secret}
            if "v2.0" in self.token_url:
                data["scope"] = self.scope
            else:
                data["resource"] = self.scope.replace("/.default", "")
            r = requests.post(self.token_url, data=data, timeout=30, verify=VERIFY_SSL, proxies=PROXIES)
            if r.status_code != 200:
                raise M365Error(f"Token request failed ({r.status_code}): {r.text[:500]}", r.status_code)
            body = r.json()
            self._token = body["access_token"]
            self._expires = time.time() + int(body.get("expires_in", 3600))
            return self._token

    @staticmethod
    def _error(r):
        code, msg = "", r.text[:500]
        try:
            err = r.json().get("error", {})
            if isinstance(err, dict):
                code, msg = err.get("code", ""), err.get("message", msg)
        except ValueError:
            pass
        return M365Error(f"HTTP {r.status_code}: {code} {msg}".strip(), r.status_code, code)

    def _sleep(self, s):
        for _ in range(max(1, int(s * 10))):
            if self.stop is not None and self.stop.is_set():
                raise M365Error("Cancelled")
            time.sleep(0.1)

    def _request(self, method, url, params=None, retries=4):
        params = dict(params or {})
        if "PublisherIdentifier" not in url:
            params.setdefault("PublisherIdentifier", self.tenant_id)
        last = None
        for attempt in range(1, retries + 1):
            self.limiter.acquire(self.stop)
            headers = {"Authorization": f"Bearer {self.token()}"}
            try:
                r = requests.request(method, url, params=params, headers=headers, timeout=120,
                                     verify=VERIFY_SSL, proxies=PROXIES)
            except requests.RequestException as e:
                last = M365Error(str(e))
                self._sleep(3 * attempt)
                continue
            if r.status_code in (200, 201):
                return r
            last = self._error(r)
            if r.status_code == 401:
                self._token = None
            elif r.status_code == 429 or r.status_code >= 500:
                wait = int(r.headers.get("Retry-After", 10 * attempt))
                if self.log:
                    self.log(f"API returned {r.status_code} - waiting {wait}s and retrying...")
                self._sleep(wait)
            else:
                raise last
        raise last

    # ---- subscriptions ----
    def list_subscriptions(self):
        r = self._request("GET", f"{self.feed}/subscriptions/list")
        return r.json()

    def start_subscription(self, ct):
        r = self._request("POST", f"{self.feed}/subscriptions/start", {"contentType": ct})
        if self.on_call:
            self.on_call("subscription_start", ct, "", "")
        return r.json()

    def stop_subscription(self, ct):
        self._request("POST", f"{self.feed}/subscriptions/stop", {"contentType": ct})
        if self.on_call:
            self.on_call("subscription_stop", ct, "", "")

    # ---- content ----
    def list_content(self, ct, start, end):
        """Yields content blobs created between start and end (max 24 h apart)."""
        url = f"{self.feed}/subscriptions/content"
        params = {"contentType": ct, "startTime": start.strftime("%Y-%m-%dT%H:%M:%S"),
                  "endTime": end.strftime("%Y-%m-%dT%H:%M:%S")}
        while url:
            r = self._request("GET", url, params)
            for item in r.json() or []:
                yield item
            url = r.headers.get("NextPageUri") or r.headers.get("NextPageUrl")
            params = None

    def fetch_blob(self, uri):
        r = self._request("GET", uri)
        data = r.json()
        return data if isinstance(data, list) else []


# ==========================================================================
#                              HELPERS
# ==========================================================================
def iso(dt):
    return dt.strftime("%Y-%m-%dT%H:%M:%SZ")


def plain(dt):
    return dt.strftime("%Y-%m-%dT%H:%M:%S")


def stamp(start, end):
    return f"{start.strftime('%Y%m%d_%H%M%S')}_{end.strftime('%Y%m%d_%H%M%S')}"


def safe_name(s):
    return re.sub(r"[^A-Za-z0-9_.-]", "_", s)


def parse_utc(s):
    s = s.strip().replace(" ", "T").replace("Z", "+00:00")
    dt = datetime.fromisoformat(s)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def disp_time(v, tz):
    if not isinstance(v, str):
        return v
    m = ISO_RE.match(v)
    if not m:
        return v
    dt = datetime.strptime(f"{m.group(1)} {m.group(2)}", "%Y-%m-%d %H:%M:%S")
    if tz == "IST":
        dt += IST
    return dt.strftime("%Y-%m-%d %H:%M:%S")


def cell_text(v, tz):
    if v is None:
        return ""
    s = json.dumps(v) if isinstance(v, (dict, list)) else str(disp_time(v, tz))
    return s.replace("\r", " ").replace("\n", " ")[:300]


def ip_of(rec):
    return rec.get("ClientIP") or rec.get("ClientIPAddress") or rec.get("ActorIpAddress") or ""


def split_terms(s):
    return [t.strip() for t in s.split(",") if t.strip()]


def clean_ip(v):
    """Normalise a ClientIP value ('1.2.3.4:443', '[2001:db8::1]:443', '::ffff:1.2.3.4') to a plain IP, or ''."""
    s = str(v or "").strip()
    if not s:
        return ""
    if s.startswith("["):
        s = s[1:].split("]")[0]
    elif s.count(":") == 1:
        s = s.split(":")[0]
    try:
        ip = ipaddress.ip_address(s)
    except ValueError:
        return ""
    if ip.version == 6 and ip.ipv4_mapped:
        ip = ip.ipv4_mapped
    return str(ip)


def is_public_ip(ip):
    try:
        return ipaddress.ip_address(ip).is_global
    except ValueError:
        return False


def ev_time(rec):
    """CreationTime as a naive UTC datetime, or None."""
    m = ISO_RE.match(str(rec.get("CreationTime", "")))
    if not m:
        return None
    return datetime.strptime(f"{m.group(1)} {m.group(2)}", "%Y-%m-%d %H:%M:%S")


# ==========================================================================
#                              DETECTIONS
# ==========================================================================
LOGIN_FAILED = {"userloginfailed"}
LOGIN_OK = {"userloggedin"}
DOWNLOAD_OPS = {"filedownloaded", "filesyncdownloadedfull"}
DELETE_OPS = {"filedeleted", "filerecycled", "filedeletedfirststagerecyclebin",
              "filedeletedsecondstagerecyclebin", "folderdeleted", "folderrecycled"}
SEVERITY_ORDER = {"High": 0, "Medium": 1, "Low": 2}


def _op(rec):
    return str(rec.get("Operation", "")).lower()


def best_window(items, minutes, distinct=False):
    """items: list of (time, value) sorted by time. Slides a window of `minutes` and returns
    (best_score, first_time, last_time, values_in_best_window). Score = event count, or the number of
    distinct values when distinct=True."""
    span = timedelta(minutes=minutes)
    best = (0, None, None, [])
    counts = Counter()
    lo = 0
    for hi, (t, v) in enumerate(items):
        counts[v] += 1
        while items[hi][0] - items[lo][0] > span:
            lv = items[lo][1]
            counts[lv] -= 1
            if counts[lv] == 0:
                del counts[lv]
            lo += 1
        score = len(counts) if distinct else hi - lo + 1
        if score > best[0]:
            best = (score, items[lo][0], t, list(counts))
    return best


def run_detections(events, cfg):
    """Threshold detections over fetched events. cfg keys: spray_users, spray_min, brute_fails, brute_min,
    dl_files, del_files, mass_min. Returns a list of hit dicts, most severe first."""
    hits = []
    fmt = lambda t: t.strftime("%Y-%m-%dT%H:%M:%S") if t else ""
    fails_by_ip, fails_by_user, ok_by_user = {}, {}, {}
    ok_pairs = set()                       # (user, ip) with a successful sign-in
    dl_by_user, del_by_user = {}, {}
    for rec in events:
        t = ev_time(rec)
        if t is None:
            continue
        op, user, ip = _op(rec), str(rec.get("UserId") or "").lower(), clean_ip(ip_of(rec))
        if op in LOGIN_FAILED:
            if ip:
                fails_by_ip.setdefault(ip, []).append((t, user))
            if user:
                fails_by_user.setdefault(user, []).append((t, ip))
        elif op in LOGIN_OK and user:
            ok_by_user.setdefault(user, []).append((t, ip))
            ok_pairs.add((user, ip))
        elif op in DOWNLOAD_OPS and user:
            dl_by_user.setdefault(user, []).append((t, rec.get("ObjectId") or rec.get("SourceFileName") or ""))
        elif op in DELETE_OPS and user:
            del_by_user.setdefault(user, []).append((t, rec.get("ObjectId") or rec.get("SourceFileName") or ""))

    # 1) Password spray: one IP failing for many different users
    for ip, items in fails_by_ip.items():
        items.sort()
        n, t0, t1, users = best_window(items, cfg["spray_min"], distinct=True)
        if n >= cfg["spray_users"]:
            hit_users = [u for u in users if (u, ip) in ok_pairs]
            hits.append({
                "Severity": "High" if hit_users else "Medium",
                "Detection": "Password spray", "MITRE": "T1110.003",
                "Entity type": "IP", "Entity": ip, "ClientIP": ip, "Count": n,
                "First seen": fmt(t0), "Last seen": fmt(t1),
                "Details": (f"Failed sign-ins for {n} different users within {cfg['spray_min']} min"
                            + (f"; {len(hit_users)} of them later signed in OK from this IP: "
                               + ", ".join(hit_users[:5]) if hit_users else "")),
                "_ops": ["UserLoginFailed", "UserLoggedIn"], "_field": "ip"})

    # 2) Brute force: many failures for one user, then (optionally) a success
    for user, items in fails_by_user.items():
        items.sort()
        n, t0, t1, ips = best_window(items, cfg["brute_min"])
        if n < cfg["brute_fails"]:
            continue
        later_ok = [x for x in ok_by_user.get(user, []) if t0 <= x[0] <= t1 + timedelta(minutes=cfg["brute_min"])]
        hits.append({
            "Severity": "High" if later_ok else "Medium",
            "Detection": "Brute force then successful sign-in" if later_ok else "Repeated failed sign-ins",
            "MITRE": "T1110.001", "Entity type": "User", "Entity": user, "UserId": user, "Count": n,
            "First seen": fmt(t0), "Last seen": fmt(later_ok[-1][0] if later_ok else t1),
            "Details": (f"{n} failed sign-ins within {cfg['brute_min']} min from "
                        f"{len({i for i in ips if i})} IP(s)"
                        + (f"; then a successful sign-in from {later_ok[0][1] or 'unknown IP'}" if later_ok else "")),
            "_ops": ["UserLoginFailed", "UserLoggedIn"], "_field": "user"})

    # 3) + 4) Mass download / mass deletion per user
    for name, data, limit, ops, mitre in (
            ("Mass file download", dl_by_user, cfg["dl_files"], ["FileDownloaded", "FileSyncDownloadedFull"],
             "T1530"),
            ("Mass file / folder deletion", del_by_user, cfg["del_files"], sorted(DELETE_OPS), "T1485")):
        for user, items in data.items():
            items.sort()
            n, t0, t1, _ = best_window(items, cfg["mass_min"])
            if n >= limit:
                hits.append({
                    "Severity": "High" if n >= limit * 3 else "Medium",
                    "Detection": name, "MITRE": mitre, "Entity type": "User", "Entity": user, "UserId": user,
                    "Count": n, "First seen": fmt(t0), "Last seen": fmt(t1),
                    "Details": f"{n} files within {cfg['mass_min']} min ({len(items)} in the whole window)",
                    "_ops": ops, "_field": "user"})

    hits.sort(key=lambda h: (SEVERITY_ORDER.get(h["Severity"], 9), -h["Count"]))
    return hits


DETECTION_COLS = ["Severity", "Detection", "MITRE", "Entity type", "Entity", "Count",
                  "First seen", "Last seen", "Details"]


class EventFilter:
    """Local filter for audit events. All conditions must match (AND)."""

    def __init__(self, workload="", operations=(), user="", ip="", text=(), result="", country="", geo=None):
        self.workload = workload.lower()
        self.ops = {o.lower() for o in operations}
        self.user, self.ip, self.result = user.lower(), ip.lower(), result.lower()
        self.text = [t.lower() for t in text]
        self.country = country.lower()
        self.geo = geo or {}             # ip -> enrichment dict from the IP Intel tab

    def active(self):
        return bool(self.workload or self.ops or self.user or self.ip or self.text or self.result or self.country)

    def describe(self):
        parts = []
        if self.workload:
            parts.append(f"workload={self.workload}")
        if self.ops:
            parts.append(f"operations={sorted(self.ops)}")
        if self.user:
            parts.append(f"user~{self.user}")
        if self.ip:
            parts.append(f"ip~{self.ip}")
        if self.result:
            parts.append(f"result~{self.result}")
        if self.text:
            parts.append(f"text any of {self.text}")
        if self.country:
            parts.append(f"country={self.country}")
        return "; ".join(parts) or "none"

    def match(self, rec):
        if self.workload and str(rec.get("Workload", "")).lower() != self.workload:
            return False
        if self.ops and str(rec.get("Operation", "")).lower() not in self.ops:
            return False
        if self.user and self.user not in str(rec.get("UserId", "")).lower():
            return False
        if self.ip and self.ip not in str(ip_of(rec)).lower():
            return False
        if self.result and self.result not in str(rec.get("ResultStatus", "")).lower():
            return False
        if self.country:
            info = self.geo.get(clean_ip(ip_of(rec)), {})
            if self.country not in (str(info.get("Country", "")).lower(), str(info.get("CC", "")).lower()):
                return False
        if self.text:
            blob = json.dumps(rec, default=str).lower()
            if not any(t in blob for t in self.text):
                return False
        return True


def order_columns(records):
    counts = Counter()
    for r in records:
        counts.update(r.keys())
    first = [c for c in PREFERRED_COLS if c in counts]
    rest = sorted((c for c in counts if c not in first), key=lambda c: (-counts[c], c))
    return first + rest


def record_row(rec, cols):
    return [json.dumps(rec, default=str) if c == "AuditData" else rec.get(c) for c in cols]


# ---------------- output sinks ----------------
class FileSink:
    def __init__(self, path, fmt):
        self.path, self.fmt = path, fmt
        self.count = 0
        self._csv = None
        self._first = True
        self.f = open(path, "w", newline="", encoding="utf-8")
        if fmt == "json":
            self.f.write("[\n")

    def write_records(self, records, cols):
        """json/jsonl write full records; csv writes the given columns."""
        for rec in records:
            if self.fmt == "jsonl":
                self.f.write(json.dumps(rec, default=str) + "\n")
            elif self.fmt == "json":
                if not self._first:
                    self.f.write(",\n")
                self.f.write(json.dumps(rec, default=str))
                self._first = False
            else:
                if self._csv is None:
                    self._csv = csv.writer(self.f)
                    self._csv.writerow(cols)
                self._csv.writerow([json.dumps(v, default=str) if isinstance(v, (dict, list)) else v
                                    for v in record_row(rec, cols)])
            self.count += 1

    def close(self):
        if self.f.closed:
            return
        if self.fmt == "json":
            self.f.write("\n]\n")
        self.f.close()


class XlsxBook:
    def __init__(self, path):
        try:
            from openpyxl import Workbook
            from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        except ImportError:
            raise M365Error("Excel output needs openpyxl. Install it with:  pip install openpyxl")
        self.path = path
        self.wb = Workbook(write_only=True)
        self.summary = self.wb.create_sheet("Summary")
        self.names = {"summary"}
        self.illegal = ILLEGAL_CHARACTERS_RE

    def clean(self, v):
        if v is None or isinstance(v, (bool, int, float)):
            return v
        s = json.dumps(v, default=str) if isinstance(v, (dict, list)) else str(v)
        return self.illegal.sub("", s)[:32000]

    def sheet(self, name):
        base = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet"
        n, i = base, 2
        while n.lower() in self.names:
            n = base[:28] + f"_{i}"
            i += 1
        self.names.add(n.lower())
        return XlsxSheet(self.wb.create_sheet(n), self, n)

    def close(self, meta, items):
        ws = self.summary
        ws.append([f"{TOOL_NAME} - export summary"])
        ws.append([])
        for k, v in meta.items():
            if isinstance(v, dict):
                v = ", ".join(f"{a}: {b}" for a, b in v.items())
            elif isinstance(v, list):
                v = ", ".join(map(str, v))
            ws.append([k, self.clean(v)])
        ws.append([])
        ws.append(["Sheet", "Rows", "Note"])
        for it in items:
            ws.append([it.get("sheet"), it["rows"], it.get("note", "")])
        self.wb.save(self.path)


class XlsxSheet:
    MAX = 1048575

    def __init__(self, ws, book, name):
        self.ws, self.book, self.name = ws, book, name
        self.count, self.header, self.truncated, self.path = 0, False, False, None

    def write_records(self, records, cols):
        if not self.header:
            self.ws.append(list(cols))
            self.header = True
        for rec in records:
            if self.count >= self.MAX:
                self.truncated = True
                return
            self.ws.append([self.book.clean(v) for v in record_row(rec, cols)])
            self.count += 1

    def close(self):
        pass


class OutputBatch:
    def __init__(self, outdir, batch_name, fmt, zip_it, meta):
        base, i = batch_name, 2
        while any(os.path.exists(os.path.join(outdir, base + ext))
                  for ext in (f".{fmt}", ".xlsx", ".meta.json", ".zip")):
            base, i = f"{batch_name}_{i}", i + 1          # never overwrite an earlier download
        self.outdir, self.batch_name, self.fmt, self.zip_it, self.meta = outdir, base, fmt, zip_it, meta
        self.items, self.files = [], []
        self.book = XlsxBook(os.path.join(outdir, base + ".xlsx")) if fmt == "xlsx" else None

    def sink(self, name):
        if self.book:
            return self.book.sheet(name)
        return FileSink(os.path.join(self.outdir, f"{self.batch_name}.{self.fmt}"), self.fmt)

    def add(self, name, sink):
        sink.close()
        item = {"name": name, "rows": sink.count}
        if self.book:
            item["sheet"] = sink.name
            if sink.truncated:
                item["note"] = "truncated at the Excel row limit"
        elif sink.count == 0:
            os.remove(sink.path)
            item["note"] = "no rows - no file created"
        else:
            item["file"] = os.path.basename(sink.path)
            self.files.append(sink.path)
        self.items.append(item)

    def finish(self):
        if self.book:
            self.book.close(self.meta, self.items)
            self.files.append(self.book.path)
        if not self.files:
            return []
        meta = dict(self.meta)
        meta["outputs"] = self.items
        mpath = os.path.join(self.outdir, self.batch_name + ".meta.json")
        with open(mpath, "w", encoding="utf-8") as f:
            json.dump(meta, f, indent=2)
        self.files.append(mpath)
        if self.zip_it:
            zpath = os.path.join(self.outdir, self.batch_name + ".zip")
            with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
                for p in self.files:
                    z.write(p, os.path.basename(p))
            for p in self.files:
                os.remove(p)
            return [zpath]
        return self.files


class ActivityLog:
    FIELDS = ["time_utc", "analyst", "machine", "tenant_id", "action", "detail", "rows", "status"]

    def __init__(self, path):
        self.path = path
        self.lock = threading.Lock()
        try:
            self.analyst = getpass.getuser()
        except Exception:
            self.analyst = "unknown"
        self.machine = platform.node()

    def write(self, action, detail="", rows="", status="ok", tenant=""):
        try:
            with self.lock:
                new = not os.path.exists(self.path)
                with open(self.path, "a", newline="", encoding="utf-8") as f:
                    w = csv.writer(f)
                    if new:
                        w.writerow(self.FIELDS)
                    w.writerow([iso(datetime.now(timezone.utc)), self.analyst, self.machine, tenant,
                                action, str(detail)[:4000], rows, status])
        except OSError:
            pass


class JsonStore:
    def __init__(self, path, default):
        self.path, self.default = path, default

    def load(self):
        try:
            with open(self.path, encoding="utf-8") as f:
                return json.load(f)
        except (OSError, ValueError):
            return self.default()

    def save(self, data):
        with open(self.path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2)


PROFILE_KEYS = ("tenant_id", "client_id", "base_url", "token_url")   # never the secret


def fetch_events(client, content_types, start, end, late_hours, stop, log, on_events, progress=None):
    """Lists content blobs for the window (plus late-arrival margin), downloads them in parallel,
    keeps events whose CreationTime is inside [start, end) and passes them to on_events(list).
    Returns (blob_count, missing_subscriptions, failed_blob_count)."""
    now = datetime.now(timezone.utc).replace(microsecond=0)
    list_start = max(start, now - RETENTION + timedelta(minutes=5))
    list_end = min(end + timedelta(hours=late_hours), now)
    if list_start >= list_end:
        log("Nothing to list: the window is older than the 7-day API retention.")
        return 0, [], 0
    blobs, missing = [], []
    for ct in content_types:
        cur, n_before = list_start, len(blobs)
        try:
            while cur < list_end:
                nxt = min(cur + timedelta(hours=24), list_end)
                for item in client.list_content(ct, cur, nxt):
                    if item.get("contentUri"):
                        blobs.append((ct, item["contentUri"]))
                cur = nxt
                if stop.is_set():
                    return 0, missing, 0
        except M365Error as e:
            if e.code == "AF20022" or "no subscription" in str(e).lower():
                missing.append(ct)
                log(f"  {ct}: no active subscription - start it in the Subscriptions panel")
                continue
            raise
        log(f"  {ct}: {len(blobs) - n_before} content blob(s)")
    if not blobs:
        return 0, missing, 0
    log(f"Downloading {len(blobs)} content blob(s) with {DOWNLOAD_THREADS} threads...")
    s0, s1 = plain(start), plain(end)
    seen = set()
    done_n = failed = 0
    pool = ThreadPoolExecutor(max_workers=DOWNLOAD_THREADS)
    try:
        futures = {pool.submit(client.fetch_blob, uri): ct for ct, uri in blobs}
        for fut in as_completed(futures):
            if stop.is_set():
                break
            done_n += 1
            try:
                events = fut.result()
            except M365Error as e:
                failed += 1
                log(f"  blob failed ({futures[fut]}): {e}")
                continue
            keep = []
            for ev in events:
                ctime = str(ev.get("CreationTime", ""))[:19]
                eid = ev.get("Id")
                if not (s0 <= ctime < s1):
                    continue
                if eid:
                    if eid in seen:
                        continue
                    seen.add(eid)
                keep.append(ev)
            if keep and on_events(keep) is False:
                stop.set()
                break
            if progress and (done_n % 10 == 0 or done_n == len(blobs)):
                progress(done_n, len(blobs))
    finally:
        pool.shutdown(wait=False, cancel_futures=True)
    if failed:
        log(f"WARNING: {failed} of {len(blobs)} content blob(s) failed to download - results are incomplete.")
    return len(blobs), missing, failed


# ==========================================================================
#                              RESULT GRID
# ==========================================================================
class ResultGrid(ttk.Frame):
    def __init__(self, parent, app, on_open=None):
        super().__init__(parent)
        self.app = app
        self.on_open = on_open          # callback(record) on double-click; default shows the record
        self.cols, self.items, self.vis = [], [], []      # items: list of (row, record)
        self.label = ""
        self.sort_col, self.sort_rev = None, False
        self.info = ttk.Label(self, text="", foreground="#0b5394")
        self.info.pack(anchor="w", pady=2)
        f = ttk.Frame(self)
        f.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(f, show="headings")
        ys = ttk.Scrollbar(f, orient="vertical", command=self.tree.yview)
        xs = ttk.Scrollbar(f, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        ys.grid(row=0, column=1, sticky="ns")
        xs.grid(row=1, column=0, sticky="ew")
        f.rowconfigure(0, weight=1)
        f.columnconfigure(0, weight=1)
        self.tree.bind("<Double-1>", lambda e: self._open())
        self.tree.bind("<Button-3>", self._right_click)
        self.tree.bind("<Button-2>", self._right_click)
        app.grids.append(self)

    def set_data(self, cols, rows, label, records=None):
        self.cols = list(cols)
        recs = records if records is not None else [dict(zip(cols, r)) for r in rows]
        self.items = list(zip(rows, recs))
        self.label = label
        self.sort_col, self.sort_rev = None, False
        self.render()

    def record(self, idx):
        return self.items[idx][1]

    def _visible(self):
        idx = list(range(len(self.cols)))
        if self.app.hide_empty.get() and self.items:
            idx = [i for i in idx if any(r[i] not in (None, "") and r[i] != [] and r[i] != {}
                                         for r, _ in self.items)]
        return idx

    def render(self):
        t = self.tree
        t.delete(*t.get_children())
        self.vis = self._visible()
        ids = [f"c{i}" for i in self.vis]
        t["displaycolumns"] = "#all"
        t["columns"] = ids
        t["displaycolumns"] = ids
        for i in self.vis:
            name = self.cols[i]
            arrow = (" \u25bc" if self.sort_rev else " \u25b2") if self.sort_col == i else ""
            t.heading(f"c{i}", text=name + arrow, command=lambda i=i: self.sort_by(i))
            t.column(f"c{i}", width=170 if name in WIDE_COLS else 120, minwidth=60, stretch=False, anchor="w")
        tz = self.app.tz_var.get()
        for n, (r, _) in enumerate(self.items[:MAX_GRID_ROWS]):
            t.insert("", "end", iid=str(n), values=[cell_text(r[i], tz) for i in self.vis])
        extra = []
        hidden = len(self.cols) - len(self.vis)
        if hidden:
            extra.append(f"{hidden} empty column(s) hidden")
        if len(self.items) > MAX_GRID_ROWS:
            extra.append(f"showing first {MAX_GRID_ROWS:,} of {len(self.items):,} - export for all")
        if self.cols:
            extra.append(f"times in {tz}")
        self.info.configure(text=self.label + ("   (" + "; ".join(extra) + ")" if extra else ""))

    def sort_by(self, i):
        if self.sort_col == i:
            self.sort_rev = not self.sort_rev
        else:
            self.sort_col, self.sort_rev = i, False

        def key(item):
            v = item[0][i]
            if isinstance(v, bool):
                return (1, str(v))
            if isinstance(v, (int, float)):
                return (0, v)
            try:
                return (0, float(v))
            except (TypeError, ValueError):
                return (1, json.dumps(v) if isinstance(v, (dict, list)) else str(v).lower())

        filled = [x for x in self.items if x[0][i] not in (None, "")]
        empty = [x for x in self.items if x[0][i] in (None, "")]
        nums = sorted([x for x in filled if key(x)[0] == 0], key=key, reverse=self.sort_rev)
        texts = sorted([x for x in filled if key(x)[0] == 1], key=key, reverse=self.sort_rev)
        self.items = (nums + texts if not self.sort_rev else texts + nums) + empty
        self.render()

    def _open(self):
        iid = self.tree.focus()
        if iid:
            (self.on_open or self.app.show_record)(self.record(int(iid)))

    def _right_click(self, event):
        row = self.tree.identify_row(event.y)
        col = self.tree.identify_column(event.x)
        if not row or not col:
            return
        pos = int(col[1:]) - 1
        if pos < 0 or pos >= len(self.vis):
            return
        ci = self.vis[pos]
        self.tree.selection_set(row)
        self.tree.focus(row)
        menu = self.app.build_cell_menu(self, int(row), self.cols[ci], self.items[int(row)][0][ci])
        menu.tk_popup(event.x_root, event.y_root)


# ==========================================================================
#                                  APP
# ==========================================================================
class App:
    CONN_FIELDS = [
        ("tenant_id", "Tenant ID", True, "Entra ID tenant GUID"),
        ("client_id", "Client ID", True, "App registration Application (client) ID"),
        ("client_secret", "Client secret", True, "Secret VALUE - kept in memory only, never saved"),
        ("base_url", "Base URL", False, f"Blank = {DEFAULT_BASE_URL} (change only for government clouds)"),
        ("token_url", "Token URL", False, "Optional. Blank = login.microsoftonline.com/<tenant>/oauth2/v2.0/token"),
    ]
    GUID_RE = re.compile(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$")

    def __init__(self, root):
        self.root = root
        self.client = None
        self.last_creds = {}
        self.q = queue.Queue()
        self.stop = threading.Event()
        self.busy = False
        self.grids = []
        self.events = []
        self.event_ids = set()
        self.filtered = []
        self.fetch_window = None
        self.fetch_types = []
        self.fetch_failed = 0
        self.fetch_note = ""
        self.window_note = ""
        self.ip_intel = {}          # ip -> enrichment dict (memory only, this session)
        self.ip_rows = []
        self._abuse_key = None      # AbuseIPDB key - memory only, never saved
        self.activity = ActivityLog(ACTIVITY_PATH)
        self.profiles = JsonStore(PROFILES_PATH, dict)
        self.tz_var = tk.StringVar(value="UTC")
        self.hide_empty = tk.BooleanVar(value=True)
        self.zip_var = tk.BooleanVar(value=False)
        self.late_var = tk.BooleanVar(value=True)
        self.ct_vars = {ct: tk.BooleanVar(value=default) for ct, _, default in CONTENT_TYPES}

        root.title(f"{TOOL_NAME} v{TOOL_VERSION}  -  developed by {DEVELOPER}, {TEAM}")
        root.geometry("1450x880")
        root.protocol("WM_DELETE_WINDOW", self.on_close)
        self._build()
        self._poll()
        self.root.after(200, lambda: self.open_connect_dialog(first=True))

    # ======================================================================
    #                                LAYOUT
    # ======================================================================
    def _build(self):
        pad = {"padx": 4, "pady": 3}
        r0 = ttk.Frame(self.root, padding=(8, 6, 8, 0))
        r0.pack(fill="x")
        self.conn_label = ttk.Label(r0, text="Not connected", foreground="#b00020")
        self.conn_label.pack(side="left", **pad)
        ttk.Button(r0, text="Disconnect", command=self.disconnect).pack(side="right", **pad)
        ttk.Button(r0, text="Connection...", command=self.open_connect_dialog).pack(side="right", **pad)

        r1 = ttk.Frame(self.root, padding=(8, 4, 8, 0))
        r1.pack(fill="x")
        ttk.Label(r1, text="Time window:").pack(side="left", **pad)
        self.preset = ttk.Combobox(r1, state="readonly", width=20, values=[p[0] for p in PRESETS])
        self.preset.current(6)
        self.preset.bind("<<ComboboxSelected>>", lambda e: self._on_preset())
        self.preset.pack(side="left", **pad)
        ttk.Label(r1, text="Minutes:").pack(side="left", **pad)
        self.minutes_var = tk.StringVar(value="1440")
        self.minutes_entry = ttk.Entry(r1, textvariable=self.minutes_var, width=8)
        self.minutes_entry.pack(side="left", **pad)
        now = datetime.now(timezone.utc)
        ttk.Label(r1, text="Start (UTC):").pack(side="left", **pad)
        self.start_var = tk.StringVar(value=(now - timedelta(days=1)).strftime("%Y-%m-%d %H:%M:%S"))
        self.start_entry = ttk.Entry(r1, textvariable=self.start_var, width=20)
        self.start_entry.pack(side="left", **pad)
        ttk.Label(r1, text="End (UTC):").pack(side="left", **pad)
        self.end_var = tk.StringVar(value=now.strftime("%Y-%m-%d %H:%M:%S"))
        self.end_entry = ttk.Entry(r1, textvariable=self.end_var, width=20)
        self.end_entry.pack(side="left", **pad)
        ttk.Label(r1, text="Audit events usually arrive 1-2 h late (up to 24 h). API keeps 7 days.",
                  foreground="gray").pack(side="left", padx=10)

        r2 = ttk.Frame(self.root, padding=(8, 2, 8, 0))
        r2.pack(fill="x")
        ttk.Label(r2, text="Logs:").pack(side="left", **pad)
        for ct, label, _ in CONTENT_TYPES:
            ttk.Checkbutton(r2, text=label, variable=self.ct_vars[ct]).pack(side="left", padx=3)
        ttk.Checkbutton(r2, text=f"Include late events (+{LATE_ARRIVAL_HOURS} h)",
                        variable=self.late_var).pack(side="left", padx=(12, 3))

        r3 = ttk.Frame(self.root, padding=(8, 2, 8, 0))
        r3.pack(fill="x")
        ttk.Button(r3, text="Fetch events", command=self.fetch).pack(side="left", padx=4)
        ttk.Button(r3, text="Download directly", command=self.download_direct).pack(side="left", padx=4)
        ttk.Button(r3, text="Cancel", command=self.cancel).pack(side="left", padx=4)
        ttk.Separator(r3, orient="vertical").pack(side="left", fill="y", padx=8)
        ttk.Label(r3, text="Format:").pack(side="left", **pad)
        self.fmt = ttk.Combobox(r3, state="readonly", width=7, values=["csv", "json", "jsonl", "xlsx"])
        self.fmt.set(DEFAULT_FORMAT)
        self.fmt.pack(side="left", **pad)
        ttk.Checkbutton(r3, text="Zip output", variable=self.zip_var).pack(side="left", padx=6)
        ttk.Label(r3, text="Show times in:").pack(side="left", **pad)
        tzc = ttk.Combobox(r3, state="readonly", width=5, values=["UTC", "IST"], textvariable=self.tz_var)
        tzc.pack(side="left", **pad)
        tzc.bind("<<ComboboxSelected>>", lambda e: self._refresh_grids())
        ttk.Checkbutton(r3, text="Hide empty columns", variable=self.hide_empty,
                        command=self._refresh_grids).pack(side="left", padx=6)

        r4 = ttk.Frame(self.root, padding=(8, 0, 8, 4))
        r4.pack(fill="x")
        ttk.Label(r4, text="Output folder:").pack(side="left", **pad)
        self.outdir_var = tk.StringVar(value=os.path.abspath(OUTPUT_DIR))
        ttk.Entry(r4, textvariable=self.outdir_var, width=70).pack(side="left", **pad)
        ttk.Button(r4, text="Browse...", command=self._choose_dir).pack(side="left", **pad)
        ttk.Button(r4, text="Open folder", command=self._open_dir).pack(side="left", **pad)

        paned = ttk.PanedWindow(self.root, orient="horizontal")
        paned.pack(fill="both", expand=True, padx=8, pady=4)
        paned.add(self._build_left(paned), weight=1)
        self.nb = ttk.Notebook(paned)
        paned.add(self.nb, weight=4)
        self._build_events_tab()
        self._build_hunts_tab()
        self._build_detections_tab()
        self._build_ip_tab()
        self._build_timeline_tab()
        self._build_log_tab()

        self.status = ttk.Label(self.root, text="Ready", relief="sunken", anchor="w", padding=(6, 2))
        self.status.pack(fill="x", side="bottom")
        self._on_preset()

    def _build_left(self, parent):
        left = ttk.PanedWindow(parent, orient="vertical")

        sub = ttk.Frame(left)
        left.add(sub, weight=1)
        ttk.Label(sub, text="Subscriptions", font=("Segoe UI", 9, "bold")).pack(anchor="w")
        self.sub_tree = ttk.Treeview(sub, columns=("ct", "status"), show="headings", height=5)
        self.sub_tree.heading("ct", text="Content type")
        self.sub_tree.heading("status", text="Status")
        self.sub_tree.column("ct", width=200)
        self.sub_tree.column("status", width=90)
        self.sub_tree.pack(fill="x", pady=2)
        sb = ttk.Frame(sub)
        sb.pack(fill="x")
        ttk.Button(sb, text="Refresh", command=self.refresh_subscriptions).pack(side="left", padx=2)
        ttk.Button(sb, text="Start selected", command=lambda: self.change_subscription(True)).pack(side="left", padx=2)
        ttk.Button(sb, text="Stop selected", command=lambda: self.change_subscription(False)).pack(side="left", padx=2)
        ttk.Label(sub, text="A content type must be started once per tenant.\n"
                            "New subscriptions can take up to 12 h to show data.",
                  foreground="gray").pack(anchor="w", pady=(2, 6))

        summ = ttk.Frame(left)
        left.add(summ, weight=3)
        top = ttk.Frame(summ)
        top.pack(fill="x")
        ttk.Label(top, text="Summary by:", font=("Segoe UI", 9, "bold")).pack(side="left")
        self.group_cb = ttk.Combobox(top, state="readonly", width=14, values=[g[0] for g in GROUP_FIELDS])
        self.group_cb.current(0)
        self.group_cb.pack(side="left", padx=4)
        self.group_cb.bind("<<ComboboxSelected>>", lambda e: self._fill_summary())
        sf = ttk.Frame(summ)
        sf.pack(fill="both", expand=True, pady=2)
        self.sum_tree = ttk.Treeview(sf, columns=("value", "count"), show="headings")
        self.sum_tree.heading("value", text="Value")
        self.sum_tree.heading("count", text="Events")
        self.sum_tree.column("value", width=230)
        self.sum_tree.column("count", width=70, anchor="e")
        ssb = ttk.Scrollbar(sf, orient="vertical", command=self.sum_tree.yview)
        self.sum_tree.configure(yscrollcommand=ssb.set)
        self.sum_tree.pack(side="left", fill="both", expand=True)
        ssb.pack(side="right", fill="y")
        self.sum_tree.bind("<Double-1>", lambda e: self._summary_drill())
        ttk.Label(summ, text="Counts follow the current filter.\nDouble-click a value to filter on it.",
                  foreground="gray").pack(anchor="w")
        return left

    def _build_events_tab(self):
        t = ttk.Frame(self.nb, padding=4)
        self.nb.add(t, text="Events")
        self.tab_events = t
        f = ttk.Frame(t)
        f.pack(fill="x")
        ttk.Label(f, text="Workload:").pack(side="left")
        self.f_workload = ttk.Combobox(f, state="readonly", width=20, values=["All"])
        self.f_workload.set("All")
        self.f_workload.pack(side="left", padx=4)
        ttk.Label(f, text="Operations:").pack(side="left")
        self.f_ops = tk.StringVar()
        e1 = ttk.Entry(f, textvariable=self.f_ops, width=40)
        e1.pack(side="left", padx=4)
        ttk.Label(f, text="Result contains:").pack(side="left")
        self.f_result = tk.StringVar()
        e2 = ttk.Entry(f, textvariable=self.f_result, width=10)
        e2.pack(side="left", padx=4)

        f2 = ttk.Frame(t)
        f2.pack(fill="x", pady=(4, 0))
        ttk.Label(f2, text="User contains:").pack(side="left")
        self.f_user = tk.StringVar()
        e3 = ttk.Entry(f2, textvariable=self.f_user, width=28)
        e3.pack(side="left", padx=4)
        ttk.Label(f2, text="Client IP contains:").pack(side="left")
        self.f_ip = tk.StringVar()
        e4 = ttk.Entry(f2, textvariable=self.f_ip, width=18)
        e4.pack(side="left", padx=4)
        ttk.Label(f2, text="Text (any of, comma-separated):").pack(side="left")
        self.f_text = tk.StringVar()
        e5 = ttk.Entry(f2, textvariable=self.f_text, width=30)
        e5.pack(side="left", padx=4)
        ttk.Label(f2, text="Country:").pack(side="left")
        self.f_country = tk.StringVar()
        e6 = ttk.Entry(f2, textvariable=self.f_country, width=14)
        e6.pack(side="left", padx=4)
        for e in (e1, e2, e3, e4, e5, e6):
            e.bind("<Return>", lambda ev: self.apply_filters())
        self.f_workload.bind("<<ComboboxSelected>>", lambda ev: self.apply_filters())

        f3 = ttk.Frame(t)
        f3.pack(fill="x", pady=4)
        ttk.Button(f3, text="Apply filter", command=self.apply_filters).pack(side="left", padx=2)
        ttk.Button(f3, text="Clear filter", command=self.clear_filters).pack(side="left", padx=2)
        ttk.Button(f3, text="Export filtered", command=lambda: self.export_events(True)).pack(side="left", padx=(12, 2))
        ttk.Button(f3, text="Export all fetched", command=lambda: self.export_events(False)).pack(side="left", padx=2)
        ttk.Label(f3, text="Operations are exact names, comma-separated (e.g. New-InboxRule, Set-Mailbox). "
                           "Right-click a cell for pivots.", foreground="gray").pack(side="left", padx=10)
        self.grid = ResultGrid(t, self)
        self.grid.pack(fill="both", expand=True)

    def _build_hunts_tab(self):
        t = ttk.Frame(self.nb, padding=4)
        self.nb.add(t, text="Quick hunts")
        self.tab_hunts = t
        pw = ttk.PanedWindow(t, orient="horizontal")
        pw.pack(fill="both", expand=True)
        lf = ttk.Frame(pw)
        pw.add(lf, weight=1)
        self.hunt_tree = ttk.Treeview(lf, show="tree")
        hsb = ttk.Scrollbar(lf, orient="vertical", command=self.hunt_tree.yview)
        self.hunt_tree.configure(yscrollcommand=hsb.set)
        self.hunt_tree.pack(side="left", fill="both", expand=True)
        hsb.pack(side="right", fill="y")
        cats = {}
        for i, h in enumerate(QUICK_HUNTS):
            c = h["category"]
            if c not in cats:
                cats[c] = self.hunt_tree.insert("", "end", text=c, open=True)
            tag = f"   [{mitre_text(h.get('mitre'))}]" if h.get("mitre") else ""
            self.hunt_tree.insert(cats[c], "end", iid=f"h{i}", text=h["name"] + tag)
        self.hunt_tree.bind("<<TreeviewSelect>>", lambda e: self._show_hunt())
        self.hunt_tree.bind("<Double-1>", lambda e: self.apply_hunt())

        rf = ttk.Frame(pw, padding=(10, 0, 0, 0))
        pw.add(rf, weight=2)
        self.hunt_name = ttk.Label(rf, text="Select a quick hunt", font=("Segoe UI", 11, "bold"))
        self.hunt_name.pack(anchor="w")
        self.hunt_desc = ttk.Label(rf, text="", wraplength=650, justify="left")
        self.hunt_desc.pack(anchor="w", pady=(4, 8))
        self.hunt_detail = tk.Text(rf, height=10, font=("Consolas", 10), wrap="word", state="disabled")
        self.hunt_detail.pack(fill="both", expand=True)
        bf = ttk.Frame(rf)
        bf.pack(fill="x", pady=6)
        ttk.Button(bf, text="Apply to fetched events", command=self.apply_hunt).pack(side="left", padx=2)
        ttk.Label(rf, text="Quick hunts set the filters on the Events tab. Fetch events for your time window "
                           "first, then apply a hunt. Use 'Summary by' (left) to group the matches, "
                           "for example by User or Client IP.", foreground="gray",
                  wraplength=650, justify="left").pack(anchor="w")

    def _build_detections_tab(self):
        t = ttk.Frame(self.nb, padding=4)
        self.nb.add(t, text="Detections")
        self.tab_detect = t
        specs = [[("spray_users", "Password spray: users per IP", SPRAY_MIN_USERS),
                  ("spray_min", "within min", SPRAY_WINDOW_MIN),
                  ("brute_fails", "Brute force: failures per user", BRUTE_MIN_FAILS),
                  ("brute_min", "within min", BRUTE_WINDOW_MIN)],
                 [("dl_files", "Mass download: files per user", MASS_DOWNLOAD_FILES),
                  ("del_files", "Mass deletion: files per user", MASS_DELETE_FILES),
                  ("mass_min", "within min", MASS_WINDOW_MIN)]]
        self.det_vars = {}
        for line in specs:
            f = ttk.Frame(t)
            f.pack(fill="x", pady=(0, 3))
            for key, label, default in line:
                ttk.Label(f, text=label + ":").pack(side="left", padx=(8, 2))
                v = tk.StringVar(value=str(default))
                ttk.Entry(f, textvariable=v, width=6).pack(side="left")
                self.det_vars[key] = v
        f = ttk.Frame(t)
        f.pack(fill="x", pady=4)
        ttk.Button(f, text="Run detections", command=self.run_detections_ui).pack(side="left", padx=2)
        ttk.Button(f, text="Export detections", command=self.export_detections).pack(side="left", padx=2)
        ttk.Label(f, text="Runs automatically after each fetch, over all fetched events (not the filter). "
                          "Double-click a row to see its events; right-click for the user timeline.",
                  foreground="gray").pack(side="left", padx=10)
        self.det_grid = ResultGrid(t, self, on_open=self._detection_drill)
        self.det_grid.pack(fill="both", expand=True)
        self.det_hits, self.det_cfg = [], {}

    def _build_ip_tab(self):
        t = ttk.Frame(self.nb, padding=4)
        self.nb.add(t, text="IP Intel")
        self.tab_ip = t
        f = ttk.Frame(t)
        f.pack(fill="x", pady=(0, 4))
        ttk.Button(f, text="Geo lookup (ip-api.com)", command=self.geo_lookup).pack(side="left", padx=2)
        ttk.Button(f, text="Check AbuseIPDB...", command=self.abuse_check).pack(side="left", padx=2)
        ttk.Button(f, text="Forget AbuseIPDB key", command=self.forget_abuse_key).pack(side="left", padx=2)
        ttk.Button(f, text="Export IP list", command=self.export_ips).pack(side="left", padx=(12, 2))
        ttk.Label(f, text="Lookups run only when you click. Results stay in memory for this session. "
                          "Double-click an IP to see its events.", foreground="gray").pack(side="left", padx=10)
        self.ip_grid = ResultGrid(t, self, on_open=lambda r: self._pivot(ip=r["IP"]))
        self.ip_grid.pack(fill="both", expand=True)

    def _build_timeline_tab(self):
        t = ttk.Frame(self.nb, padding=4)
        self.nb.add(t, text="User Timeline")
        self.tab_timeline = t
        f = ttk.Frame(t)
        f.pack(fill="x", pady=(0, 4))
        ttk.Label(f, text="User:").pack(side="left")
        self.tl_user = tk.StringVar()
        self.tl_cb = ttk.Combobox(f, textvariable=self.tl_user, width=48)
        self.tl_cb.pack(side="left", padx=4)
        self.tl_cb.bind("<Return>", lambda e: self.show_timeline())
        self.tl_cb.bind("<<ComboboxSelected>>", lambda e: self.show_timeline())
        ttk.Button(f, text="Show timeline", command=self.show_timeline).pack(side="left", padx=2)
        self.tl_flagged = tk.BooleanVar(value=False)
        ttk.Checkbutton(f, text="Flagged events only", variable=self.tl_flagged,
                        command=self.show_timeline).pack(side="left", padx=8)
        ttk.Button(f, text="Export timeline", command=self.export_timeline).pack(side="left", padx=(12, 2))
        ttk.Label(f, text="Oldest first. 'Hunt flags' shows which quick hunts an event matches.",
                  foreground="gray").pack(side="left", padx=10)
        self.tl_grid = ResultGrid(t, self)
        self.tl_grid.pack(fill="both", expand=True)
        self.tl_rows = []

    def _build_log_tab(self):
        t = ttk.Frame(self.nb, padding=4)
        self.nb.add(t, text="Log")
        self.logbox = tk.Text(t, font=("Consolas", 9), state="disabled", wrap="none")
        ls = ttk.Scrollbar(t, orient="vertical", command=self.logbox.yview)
        self.logbox.configure(yscrollcommand=ls.set)
        self.logbox.pack(side="left", fill="both", expand=True)
        ls.pack(side="right", fill="y")

    # ======================================================================
    #                              CONNECTION
    # ======================================================================
    def open_connect_dialog(self, first=False):
        if self.busy:
            messagebox.showinfo("Busy", "Wait for the current task to finish.")
            return
        dlg = tk.Toplevel(self.root)
        dlg.title("Connect to Microsoft 365 audit log")
        dlg.transient(self.root)
        dlg.resizable(False, False)
        frm = ttk.Frame(dlg, padding=14)
        frm.pack(fill="both", expand=True)
        ttk.Label(frm, text=f"{TOOL_NAME}  v{TOOL_VERSION}",
                  font=("Segoe UI", 12, "bold")).grid(row=0, column=0, columnspan=3, sticky="w")
        ttk.Label(frm, text=f"Developed by {DEVELOPER} - {TEAM}",
                  foreground="gray").grid(row=1, column=0, columnspan=3, sticky="w", pady=(0, 8))

        pf = ttk.Frame(frm)
        pf.grid(row=2, column=0, columnspan=3, sticky="we", pady=(0, 6))
        ttk.Label(pf, text="Profile:").pack(side="left")
        prof_cb = ttk.Combobox(pf, state="readonly", width=30, values=sorted(self.profiles.load()))
        prof_cb.pack(side="left", padx=6)

        vars_, entries = {}, {}
        row = 3
        for key, label, required, hint in self.CONN_FIELDS:
            ttk.Label(frm, text=label + (" *" if required else "")).grid(row=row, column=0, sticky="w", pady=(4, 0))
            v = tk.StringVar(value="" if key == "client_secret" else self.last_creds.get(key, ""))
            e = ttk.Entry(frm, textvariable=v, width=62, show="\u2022" if key == "client_secret" else "")
            e.grid(row=row, column=1, sticky="we", pady=(4, 0), padx=6)
            vars_[key], entries[key] = v, e
            ttk.Label(frm, text=hint, foreground="gray").grid(row=row + 1, column=1, sticky="w", padx=6)
            row += 2

        def apply_profile(_=None):
            p = self.profiles.load().get(prof_cb.get(), {})
            for k in PROFILE_KEYS:
                vars_[k].set(p.get(k, ""))
            vars_["client_secret"].set("")
            entries["client_secret"].focus_set()
        prof_cb.bind("<<ComboboxSelected>>", apply_profile)

        def save_profile():
            name = simpledialog.askstring("Save profile", "Profile name (e.g. client name):", parent=dlg)
            if not name or not name.strip():
                return
            data = self.profiles.load()
            data[name.strip()] = {k: vars_[k].get().strip() for k in PROFILE_KEYS}
            try:
                self.profiles.save(data)
            except OSError as e:
                messagebox.showerror("Profile", f"Could not save profile: {e}", parent=dlg)
                return
            prof_cb.configure(values=sorted(data))
            prof_cb.set(name.strip())
            status.configure(text=f"Profile '{name.strip()}' saved (without the secret).", foreground="#1e7b34")

        def delete_profile():
            name = prof_cb.get()
            if not name or not messagebox.askyesno("Delete profile", f"Delete profile '{name}'?", parent=dlg):
                return
            data = self.profiles.load()
            data.pop(name, None)
            self.profiles.save(data)
            prof_cb.configure(values=sorted(data))
            prof_cb.set("")

        ttk.Button(pf, text="Save as profile...", command=save_profile).pack(side="left", padx=2)
        ttk.Button(pf, text="Delete profile", command=delete_profile).pack(side="left", padx=2)

        show = tk.BooleanVar(value=False)
        ttk.Checkbutton(frm, text="Show secret", variable=show,
                        command=lambda: entries["client_secret"].configure(show="" if show.get() else "\u2022")
                        ).grid(row=row, column=1, sticky="w", padx=6, pady=4)
        row += 1
        ttk.Label(frm, foreground="#555",
                  text="The secret is never saved. Profiles store only tenant ID, client ID and URLs."
                  ).grid(row=row, column=0, columnspan=3, sticky="w", pady=(6, 0))
        row += 1
        status = ttk.Label(frm, text="", foreground="#b00020", wraplength=600)
        status.grid(row=row, column=0, columnspan=3, sticky="w", pady=6)
        row += 1
        btns = ttk.Frame(frm)
        btns.grid(row=row, column=0, columnspan=3, sticky="e")
        connect_btn = ttk.Button(btns, text="Connect")
        connect_btn.pack(side="left", padx=4)

        def cancel():
            dlg.grab_release()
            dlg.destroy()
            if first and self.client is None:
                self.on_close(confirm=False)

        ttk.Button(btns, text="Exit" if first else "Cancel", command=cancel).pack(side="left", padx=4)

        def connect(*_):
            creds = {k: v.get().strip() for k, v in vars_.items()}
            missing = [lbl for k, lbl, req, _ in self.CONN_FIELDS if req and not creds[k]]
            if missing:
                status.configure(text="Required: " + ", ".join(missing), foreground="#b00020")
                return
            for k in ("tenant_id", "client_id"):
                if not self.GUID_RE.match(creds[k]):
                    lbl = dict((f[0], f[1]) for f in self.CONN_FIELDS)[k]
                    status.configure(text=f"{lbl} must be a GUID like xxxxxxxx-xxxx-xxxx-xxxx-xxxxxxxxxxxx",
                                     foreground="#b00020")
                    return
            if creds["base_url"] and not creds["base_url"].lower().startswith("https://"):
                status.configure(text="Base URL must start with https://", foreground="#b00020")
                return
            client = M365Client(creds)
            status.configure(text="Connecting...", foreground="#0b5394")
            connect_btn.configure(state="disabled")

            def fn():
                try:
                    client.token()
                    return ("ok", client.list_subscriptions())
                except M365Error as e:
                    if e.status in (401, 403):
                        return ("err", f"{e}\n\nCheck that the app has the API permission 'Office 365 Management "
                                       "APIs' -> Application -> ActivityFeed.Read with admin consent.")
                    return ("err", str(e))
                except Exception as e:
                    return ("err", str(e))

            def done(res):
                if not dlg.winfo_exists():
                    return
                connect_btn.configure(state="normal")
                if res[0] != "ok":
                    client.wipe()
                    self.activity.write("connect", client.feed, status="failed", tenant=creds["tenant_id"])
                    status.configure(text=f"Connection failed: {res[1]}", foreground="#b00020")
                    return
                if self.client:
                    self.client.wipe()
                self.client = client
                client.log = self.log
                client.stop = self.stop
                tenant = client.tenant_id
                client.on_call = lambda action, detail, n, err: self.activity.write(
                    action, detail, n, f"error: {err}" if err else "ok", tenant)
                self.last_creds = {k: v for k, v in creds.items() if k != "client_secret"}
                name = prof_cb.get()
                self.conn_label.configure(
                    text=f"Connected  |  {('Profile ' + name + '  |  ') if name else ''}Tenant {tenant}",
                    foreground="#1e7b34")
                self.activity.write("connect", client.feed, tenant=tenant)
                self._fill_subscriptions(res[1])
                self.log(f"Connected. Feed: {client.feed}")
                self.log(f"Activity log: {ACTIVITY_PATH}")
                dlg.grab_release()
                dlg.destroy()
            self.run_bg(fn, done, "Connecting...")

        connect_btn.configure(command=connect)
        dlg.bind("<Return>", connect)
        dlg.protocol("WM_DELETE_WINDOW", cancel)
        dlg.update_idletasks()
        x = self.root.winfo_rootx() + (self.root.winfo_width() - dlg.winfo_width()) // 2
        y = self.root.winfo_rooty() + (self.root.winfo_height() - dlg.winfo_height()) // 3
        dlg.geometry(f"+{max(x, 0)}+{max(y, 0)}")
        dlg.grab_set()
        first_empty = next((k for k, *_ in self.CONN_FIELDS if not vars_[k].get()), "tenant_id")
        entries[first_empty].focus_set()

    def disconnect(self):
        if self.busy:
            messagebox.showinfo("Busy", "Wait for the current task to finish or click Cancel.")
            return
        if self.client:
            self.activity.write("disconnect", tenant=self._tenant())
            self.client.wipe()
        self.client = None
        self._abuse_key = None
        self.sub_tree.delete(*self.sub_tree.get_children())
        self.conn_label.configure(text="Not connected", foreground="#b00020")
        self.log("Disconnected. Credentials cleared from memory.")

    def _need_client(self):
        if self.client is None:
            messagebox.showinfo("Not connected", "Connect first.")
            self.open_connect_dialog()
            return False
        return True

    def on_close(self, confirm=True):
        if confirm and self.busy and not messagebox.askyesno("Exit", "A task is still running. Exit anyway?"):
            return
        if self.client:
            self.activity.write("close", tenant=self._tenant())
            self.client.wipe()
        self.client = None
        self.last_creds = {}
        self._abuse_key = None
        self.root.destroy()

    # ======================================================================
    #                              PLUMBING
    # ======================================================================
    def log(self, msg):
        line = f"{datetime.now().strftime('%H:%M:%S')}  {msg}"
        if threading.current_thread() is threading.main_thread():
            self._append_log(line)
        else:
            self.q.put(("log", None, line))

    def ui(self, fn):
        self.q.put(("ui", fn, None))

    def _append_log(self, line):
        self.logbox.configure(state="normal")
        self.logbox.insert("end", line + "\n")
        self.logbox.see("end")
        self.logbox.configure(state="disabled")
        self.status.configure(text=line)

    def _poll(self):
        try:
            while True:
                kind, cb, payload = self.q.get_nowait()
                if kind == "log":
                    self._append_log(payload)
                elif kind == "ui":
                    cb()
                elif kind == "done":
                    self._set_busy(False)
                    if cb:
                        cb(payload)
                elif kind == "error":
                    self._set_busy(False)
                    self.log(f"ERROR: {payload}")
                    messagebox.showerror("Error", str(payload))
        except queue.Empty:
            pass
        self.root.after(100, self._poll)

    def _set_busy(self, busy, msg=""):
        self.busy = busy
        self.root.configure(cursor="watch" if busy else "")
        self.status.configure(text=msg if busy else "Ready")

    def run_bg(self, fn, done=None, msg="Working..."):
        if self.busy:
            messagebox.showinfo("Busy", "Another task is running. Wait or click Cancel.")
            return
        self.stop.clear()
        self._set_busy(True, msg)
        self.log(msg)

        def worker():
            try:
                self.q.put(("done", done, fn()))
            except Exception as e:
                self.q.put(("error", None, e))
        threading.Thread(target=worker, daemon=True).start()

    def cancel(self):
        if self.busy:
            self.stop.set()
            self.log("Cancel requested...")

    def _refresh_grids(self):
        for g in self.grids:
            g.render()

    # ======================================================================
    #                              INPUTS
    # ======================================================================
    def _on_preset(self):
        val = PRESETS[self.preset.current()][1]
        self.minutes_entry.configure(state="normal" if val == "minutes" else "disabled")
        st = "normal" if val == "range" else "disabled"
        self.start_entry.configure(state=st)
        self.end_entry.configure(state=st)
        if isinstance(val, int):
            self.minutes_var.set(str(val))

    def window(self):
        val = PRESETS[self.preset.current()][1]
        now = datetime.now(timezone.utc).replace(microsecond=0)
        if val == "range":
            start, end = parse_utc(self.start_var.get()), parse_utc(self.end_var.get())
        else:
            mins = int(self.minutes_var.get()) if val == "minutes" else val
            start, end = now - timedelta(minutes=mins), now
        if start >= end:
            raise ValueError("Start must be before end")
        if end <= now - RETENTION:
            raise ValueError("The whole window is older than 7 days - the API no longer has this data")
        self.window_note = ""
        if start < now - RETENTION:
            kept = now - RETENTION + timedelta(minutes=5)
            self.window_note = (f"Window clipped to the 7-day API retention: requested start {iso(start)}, "
                                f"data only available from about {iso(kept)}.")
        return start, end

    def _get_window_or_warn(self):
        try:
            w = self.window()
        except ValueError as e:
            messagebox.showerror("Time window", f"Invalid time window: {e}\nUse format YYYY-MM-DD HH:MM:SS (UTC).")
            return None
        if self.window_note:
            self.log(self.window_note)
            if not messagebox.askokcancel("Time window clipped", self.window_note +
                                          "\n\nThe older part cannot be fetched. Continue with what is available?"):
                return None
        return w

    def _content_types(self):
        cts = [ct for ct, _, _ in CONTENT_TYPES if self.ct_vars[ct].get()]
        if not cts:
            messagebox.showinfo("Logs", "Tick at least one log type.")
        return cts

    def _outdir(self):
        d = self.outdir_var.get().strip() or OUTPUT_DIR
        os.makedirs(d, exist_ok=True)
        return d

    def _choose_dir(self):
        d = filedialog.askdirectory(initialdir=self.outdir_var.get())
        if d:
            self.outdir_var.set(d)

    def _open_dir(self):
        d = self._outdir()
        if sys.platform.startswith("win"):
            os.startfile(d)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", d])
        else:
            subprocess.Popen(["xdg-open", d])

    def _meta(self, start, end, fmt, cts, flt=None):
        return {
            "tool": f"{TOOL_NAME} v{TOOL_VERSION}",
            "developed_by": f"{DEVELOPER}, {TEAM}",
            "tenant_id": self.client.tenant_id if self.client else "",
            "analyst": self.activity.analyst,
            "machine": self.activity.machine,
            "created_utc": iso(datetime.now(timezone.utc)),
            "time_window_utc": {"start": iso(start), "end": iso(end)},
            "content_types": cts,
            "filter": flt.describe() if flt else "none",
            "format": fmt,
            "window_note": self.window_note or "none",
        }

    def current_filter(self):
        wl = self.f_workload.get()
        return EventFilter(workload="" if wl in ("", "All") else wl, operations=split_terms(self.f_ops.get()),
                           user=self.f_user.get().strip(), ip=self.f_ip.get().strip(),
                           text=split_terms(self.f_text.get()), result=self.f_result.get().strip(),
                           country=self.f_country.get().strip(), geo=dict(self.ip_intel))

    # ======================================================================
    #                            SUBSCRIPTIONS
    # ======================================================================
    def _fill_subscriptions(self, subs):
        status = {s.get("contentType"): s.get("status", "") for s in subs or []}
        self.sub_tree.delete(*self.sub_tree.get_children())
        for ct, _, _ in CONTENT_TYPES:
            self.sub_tree.insert("", "end", iid=ct, values=(ct, status.get(ct, "not started")))
        off = [ct for ct, _, d in CONTENT_TYPES if d and status.get(ct, "").lower() != "enabled"]
        if off:
            self.log(f"Not enabled: {', '.join(off)}. Select them under Subscriptions and click 'Start selected'.")

    def refresh_subscriptions(self):
        if not self._need_client():
            return
        self.run_bg(self.client.list_subscriptions, self._fill_subscriptions, "Reading subscriptions...")

    def change_subscription(self, start):
        if not self._need_client():
            return
        sel = list(self.sub_tree.selection())
        if not sel:
            messagebox.showinfo("Subscriptions", "Select one or more content types first.")
            return
        if not start and not messagebox.askyesno(
                "Stop subscription", "Stopping a subscription stops collection for every tool using it "
                                     "in this tenant. Continue?"):
            return

        def fn():
            for ct in sel:
                try:
                    if start:
                        self.client.start_subscription(ct)
                    else:
                        self.client.stop_subscription(ct)
                    self.log(f"{'Started' if start else 'Stopped'} {ct}")
                except M365Error as e:
                    self.log(f"{ct}: {e}")
            return self.client.list_subscriptions()

        def done(subs):
            self._fill_subscriptions(subs)
            if start:
                messagebox.showinfo("Subscriptions", "Started. Note: new subscriptions can take up to 12 hours "
                                                     "before content is available, and only cover events from "
                                                     "the start time onward.")
        self.run_bg(fn, done, "Updating subscriptions...")

    # ======================================================================
    #                                FETCH
    # ======================================================================
    def fetch(self):
        if not self._need_client():
            return
        w = self._get_window_or_warn()
        cts = self._content_types()
        if not w or not cts:
            return
        start, end = w
        late = LATE_ARRIVAL_HOURS if self.late_var.get() else 0
        self.events, self.event_ids, self.filtered = [], set(), []
        self.grid.set_data([], [], "Fetching...")
        state = {"capped": False}

        def on_events(evs):
            room = MAX_EVENTS_IN_MEMORY - len(self.events)
            for ev in evs:
                eid = ev.get("Id")
                if eid and eid in self.event_ids:
                    continue
                if room <= 0:
                    state["capped"] = True
                    return False
                if eid:
                    self.event_ids.add(eid)
                self.events.append(ev)
                room -= 1
            return True

        def progress(n, total):
            self.ui(lambda: self.status.configure(
                text=f"Downloaded {n}/{total} blobs - {len(self.events):,} events in window"))

        def fn():
            return fetch_events(self.client, cts, start, end, late, self.stop, self.log, on_events, progress)

        def done(res):
            blobs, missing, failed = res
            self.fetch_window, self.fetch_types = (start, end), cts
            self.fetch_failed, self.fetch_note = failed, self.window_note
            self.events.sort(key=lambda e: str(e.get("CreationTime", "")), reverse=True)
            wls = sorted({str(e.get("Workload", "")) for e in self.events if e.get("Workload")})
            self.f_workload.configure(values=["All"] + wls)
            if self.f_workload.get() not in ["All"] + wls:
                self.f_workload.set("All")
            msg = f"Fetched {len(self.events):,} events from {blobs} blob(s) for {iso(start)} -> {iso(end)}"
            self.log(msg)
            self.activity.write("fetch", f"{','.join(cts)} {iso(start)}->{iso(end)}", len(self.events),
                                tenant=self._tenant())
            if failed:
                messagebox.showwarning("Incomplete results",
                                       f"{failed} of {blobs} content blob(s) failed to download after retries, "
                                       "so some events are missing. See the Log tab. Fetch again to retry.")
            if state["capped"]:
                messagebox.showwarning("Event limit", f"Stopped at {MAX_EVENTS_IN_MEMORY:,} events to protect "
                                       "memory. Use a shorter window, or 'Download directly' to save everything "
                                       "to a file.")
            if missing:
                messagebox.showinfo("Subscriptions", "No active subscription for:\n" + "\n".join(missing) +
                                    "\n\nSelect them in the Subscriptions panel and click 'Start selected'.")
            elif not self.events:
                messagebox.showinfo("Fetch", "No events in this window. Audit events usually arrive 1-2 hours "
                                    "late - try a longer or earlier window.")
            self.apply_filters()
            self._build_ip_table()
            self._fill_timeline_users()
            self.run_detections_ui(auto=True)
        self.run_bg(fn, done, f"Fetching {', '.join(cts)}...")

    # ======================================================================
    #                           FILTER + SUMMARY
    # ======================================================================
    def apply_filters(self):
        flt = self.current_filter()
        self.filtered = [e for e in self.events if flt.match(e)] if flt.active() else list(self.events)
        cols = order_columns(self.filtered)
        rows = [[e.get(c) for c in cols] for e in self.filtered]
        label = (f"{len(self.filtered):,} of {len(self.events):,} fetched events"
                 + (f"   |   filter: {flt.describe()}" if flt.active() else ""))
        self.grid.set_data(cols, rows, label, self.filtered)
        self._fill_summary()
        self.nb.select(self.tab_events)

    def clear_filters(self):
        self.f_workload.set("All")
        for v in (self.f_ops, self.f_user, self.f_ip, self.f_text, self.f_result, self.f_country):
            v.set("")
        self.apply_filters()

    def _group_value(self, rec, field):
        if field == "_ip":
            return ip_of(rec)
        if field == "_country":
            info = self.ip_intel.get(clean_ip(ip_of(rec)))
            if not ip_of(rec):
                return ""
            return info.get("Country") or "(not looked up)" if info else "(not looked up)"
        v = rec.get(field)
        if field == "RecordType" and isinstance(v, int) and v in RECORD_TYPES:
            return f"{v} ({RECORD_TYPES[v]})"
        return v

    def _fill_summary(self):
        label, field = GROUP_FIELDS[self.group_cb.current()]
        counts = Counter(str(self._group_value(e, field) or "(empty)") for e in self.filtered)
        self.sum_tree.delete(*self.sum_tree.get_children())
        for val, n in counts.most_common(500):
            self.sum_tree.insert("", "end", values=(val, f"{n:,}"))

    def _summary_drill(self):
        iid = self.sum_tree.focus()
        if not iid:
            return
        val = self.sum_tree.item(iid, "values")[0]
        if val in ("(empty)", "(not looked up)"):
            if val == "(not looked up)":
                messagebox.showinfo("Country", "Run 'Geo lookup' on the IP Intel tab first.")
            return
        label, field = GROUP_FIELDS[self.group_cb.current()]
        self._filter_field(field, val)

    def _filter_field(self, field, val):
        val = str(val)
        if field == "Operation":
            self.f_ops.set(val)
        elif field == "UserId":
            self.f_user.set(val)
        elif field in ("_ip", "ClientIP", "ClientIPAddress", "ActorIpAddress"):
            self.f_ip.set(val)
        elif field == "Workload":
            self.f_workload.set(val)
        elif field == "ResultStatus":
            self.f_result.set(val)
        elif field == "_country":
            self.f_country.set(val)
        else:
            self.f_text.set(val.split(" (")[0] if field == "RecordType" else val)
        self.apply_filters()

    # ======================================================================
    #                             QUICK HUNTS
    # ======================================================================
    def _hunt(self):
        iid = self.hunt_tree.focus()
        return QUICK_HUNTS[int(iid[1:])] if iid.startswith("h") else None

    def _show_hunt(self):
        h = self._hunt()
        if not h:
            return
        self.hunt_name.configure(text=h["name"])
        self.hunt_desc.configure(text=h["description"])
        lines = []
        if h.get("mitre"):
            lines.append("MITRE ATT&CK:\n  " + mitre_text(h["mitre"], names=True).replace("\n", "\n  "))
        if h.get("operations"):
            lines.append("Operations (any of):\n  " + "\n  ".join(h["operations"]))
        if h.get("text"):
            lines.append("AND event text contains any of:\n  " + "\n  ".join(h["text"]))
        if h.get("result"):
            lines.append(f"Result status contains: {h['result']}")
        self.hunt_detail.configure(state="normal")
        self.hunt_detail.delete("1.0", "end")
        self.hunt_detail.insert("1.0", "\n\n".join(lines))
        self.hunt_detail.configure(state="disabled")

    def apply_hunt(self):
        h = self._hunt()
        if not h:
            messagebox.showinfo("Quick hunts", "Select a quick hunt first.")
            return
        if not self.events:
            messagebox.showinfo("Quick hunts", "Fetch events first (pick a time window and click 'Fetch events').")
            return
        self.f_workload.set("All")
        self.f_ops.set(", ".join(h.get("operations", [])))
        self.f_text.set(", ".join(h.get("text", [])))
        self.f_result.set(h.get("result", ""))
        self.f_user.set("")
        self.f_ip.set("")
        self.f_country.set("")
        self.log(f"Quick hunt: {h['name']}")
        self.apply_filters()

    # ======================================================================
    #                          RECORD + PIVOTS
    # ======================================================================
    def show_record(self, rec):
        text = json.dumps(rec, indent=2, default=str)
        win = tk.Toplevel(self.root)
        win.title("Audit record")
        win.geometry("780x680")
        txt = tk.Text(win, font=("Consolas", 10), wrap="word")
        sb = ttk.Scrollbar(win, orient="vertical", command=txt.yview)
        txt.configure(yscrollcommand=sb.set)
        bar = ttk.Frame(win)
        bar.pack(side="bottom", fill="x")
        ttk.Button(bar, text="Copy JSON", command=lambda: self._copy(text)).pack(side="left", padx=4, pady=4)
        ttk.Button(bar, text="Close", command=win.destroy).pack(side="right", padx=4, pady=4)
        sb.pack(side="right", fill="y")
        txt.pack(fill="both", expand=True)
        txt.insert("1.0", text)
        txt.configure(state="disabled")

    def _copy(self, text):
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        self.log("Copied to clipboard.")

    def build_cell_menu(self, grid, rowidx, name, val):
        m = tk.Menu(self.root, tearoff=0)
        sval = "" if val is None else (json.dumps(val) if isinstance(val, (dict, list)) else str(val))
        rec = grid.record(rowidx)
        m.add_command(label="Copy value", command=lambda: self._copy(sval))
        m.add_command(label="Copy record as JSON", command=lambda: self._copy(json.dumps(rec, indent=2, default=str)))
        m.add_command(label="Show full record", command=lambda: self.show_record(rec))
        if not sval or isinstance(val, (dict, list)):
            return m
        m.add_separator()
        if name in ("Operation", "UserId", "Workload", "ResultStatus", "ClientIP", "ClientIPAddress", "ActorIpAddress"):
            m.add_command(label=f"Filter: {name} = this value", command=lambda: self._filter_field(name, val))
        else:
            m.add_command(label="Filter: events containing this value", command=lambda: self._filter_field("_text", val))
        user = rec.get("UserId")
        if user:
            m.add_command(label=f"All activity for {user}", command=lambda: self._pivot(user=str(user)))
            m.add_command(label=f"User timeline for {user}", command=lambda: self.show_timeline(str(user)))
        ip = ip_of(rec)
        if ip:
            m.add_command(label=f"All activity from {ip}", command=lambda: self._pivot(ip=str(ip)))
        return m

    def _pivot(self, user="", ip=""):
        self.f_workload.set("All")
        for v in (self.f_ops, self.f_text, self.f_result, self.f_country):
            v.set("")
        self.f_user.set(user)
        self.f_ip.set(ip)
        self.apply_filters()
        self.grid.sort_by(self.grid.cols.index("CreationTime")) if "CreationTime" in self.grid.cols else None

    # ======================================================================
    #                               EXPORT
    # ======================================================================
    def export_events(self, filtered):
        recs = self.filtered if filtered else self.events
        if not recs:
            messagebox.showinfo("Export", "Nothing to export - fetch events first.")
            return
        start, end = self.fetch_window
        fmt, outdir, zip_it = self.fmt.get(), self._outdir(), self.zip_var.get()
        flt = self.current_filter() if filtered else None
        meta = self._meta(start, end, fmt, self.fetch_types, flt)
        meta["window_note"] = self.fetch_note or "none"
        meta["failed_blobs"] = self.fetch_failed
        meta["complete"] = self.fetch_failed == 0
        name = f"m365_audit_{'filtered_' if filtered and flt.active() else ''}{stamp(start, end)}"
        recs = list(recs)

        def fn():
            cols = order_columns(recs)
            batch = OutputBatch(outdir, name, fmt, zip_it, meta)
            sink = batch.sink("AuditLog")
            sink.write_records(sorted(recs, key=lambda e: str(e.get("CreationTime", ""))), cols)
            batch.add("AuditLog", sink)
            files = batch.finish()
            self.activity.write("export", "; ".join(files), len(recs), tenant=self._tenant())
            return files

        def done(files):
            for f in files:
                self.log(f"Exported -> {f}")
            messagebox.showinfo("Export finished", f"{len(recs):,} events\n\n" +
                                "\n".join(os.path.basename(f) for f in files) + f"\n\nFolder: {outdir}")
        self.run_bg(fn, done, f"Exporting {len(recs):,} events...")

    def download_direct(self):
        """Stream everything in the window straight to a file (no memory limit). Current filter applies."""
        if not self._need_client():
            return
        w = self._get_window_or_warn()
        cts = self._content_types()
        if not w or not cts:
            return
        start, end = w
        late = LATE_ARRIVAL_HOURS if self.late_var.get() else 0
        fmt, outdir, zip_it = self.fmt.get(), self._outdir(), self.zip_var.get()
        flt = self.current_filter()
        if flt.active() and not messagebox.askyesno(
                "Download directly", f"The current filter will be applied:\n{flt.describe()}\n\nContinue?"):
            return
        meta = self._meta(start, end, fmt, cts, flt if flt.active() else None)
        if fmt in ("csv", "xlsx"):
            meta["columns_note"] = "Common fields as columns; full record in the AuditData column (JSON)."
        name = f"m365_audit_{stamp(start, end)}"

        def fn():
            batch = OutputBatch(outdir, name, fmt, zip_it, meta)
            sink = batch.sink("AuditLog")
            lock = threading.Lock()

            def on_events(evs):
                keep = [e for e in evs if flt.match(e)] if flt.active() else evs
                if keep:
                    with lock:
                        sink.write_records(keep, STREAM_COLS)
                return True

            def progress(n, total):
                self.ui(lambda: self.status.configure(
                    text=f"Downloaded {n}/{total} blobs - {sink.count:,} events written"))
            failed = 0
            try:
                blobs, missing, failed = fetch_events(self.client, cts, start, end, late, self.stop, self.log,
                                                      on_events, progress)
            finally:
                batch.add("AuditLog", sink)
                batch.meta["failed_blobs"] = failed
                batch.meta["complete"] = failed == 0
            files = batch.finish()
            self.activity.write("download", "; ".join(files), sink.count,
                                "ok" if not failed else f"incomplete: {failed} blob(s) failed",
                                tenant=self._tenant())
            return sink.count, files, missing, failed

        def done(res):
            n, files, missing, failed = res
            note = ("\n\nNo active subscription for: " + ", ".join(missing)) if missing else ""
            if failed:
                note += (f"\n\nWARNING: {failed} content blob(s) failed to download - the file is incomplete. "
                         "This is recorded in the .meta.json file.")
            self.log(f"Download finished: {n:,} events")
            for f in files:
                self.log(f"  -> {f}")
            (messagebox.showwarning if failed else messagebox.showinfo)("Download finished", f"{n:,} events\n\n" +
                                ("\n".join(os.path.basename(f) for f in files) or "No events - no file created") +
                                f"\n\nFolder: {outdir}{note}")
        self.run_bg(fn, done, f"Downloading {', '.join(cts)} directly to file...")

    # ======================================================================
    #                              DETECTIONS
    # ======================================================================
    def _tenant(self):
        return self.client.tenant_id if self.client else ""

    def _wait(self, seconds):
        for _ in range(max(1, int(seconds * 10))):
            if self.stop.is_set():
                raise M365Error("Cancelled")
            time.sleep(0.1)

    def _det_cfg(self):
        try:
            cfg = {k: int(v.get()) for k, v in self.det_vars.items()}
        except ValueError:
            raise ValueError("Detection thresholds must be whole numbers.")
        if any(v <= 0 for v in cfg.values()):
            raise ValueError("Detection thresholds must be greater than 0.")
        return cfg

    def run_detections_ui(self, auto=False):
        if not self.events:
            if not auto:
                messagebox.showinfo("Detections", "Fetch events first.")
            return
        try:
            cfg = self._det_cfg()
        except ValueError as e:
            if not auto:
                messagebox.showerror("Detections", str(e))
            return
        hits = run_detections(self.events, cfg)
        self.det_hits, self.det_cfg = hits, cfg
        high = sum(h["Severity"] == "High" for h in hits)
        rows = [[h.get(c) for c in DETECTION_COLS] for h in hits]
        label = (f"{len(hits)} detection(s): {high} high, {len(hits) - high} medium - over "
                 f"{len(self.events):,} fetched events" if hits else
                 f"No detections over {len(self.events):,} fetched events with these thresholds")
        self.det_grid.set_data(DETECTION_COLS, rows, label, hits)
        self.nb.tab(self.tab_detect, text=f"Detections ({len(hits)})" if hits else "Detections")
        self.activity.write("detections", json.dumps(cfg), len(hits), tenant=self._tenant())
        self.log(label)
        if not auto:
            self.nb.select(self.tab_detect)

    def _detection_drill(self, hit):
        self.f_workload.set("All")
        for v in (self.f_text, self.f_result, self.f_country, self.f_user, self.f_ip):
            v.set("")
        self.f_ops.set(", ".join(hit.get("_ops", [])))
        (self.f_ip if hit.get("_field") == "ip" else self.f_user).set(hit["Entity"])
        self.apply_filters()
        if "CreationTime" in self.grid.cols:
            self.grid.sort_by(self.grid.cols.index("CreationTime"))

    def export_detections(self):
        recs = [{c: h.get(c) for c in DETECTION_COLS} for h in self.det_hits]
        self._export_records("m365_detections", "Detections", DETECTION_COLS, recs,
                             {"thresholds": self.det_cfg})

    # ======================================================================
    #                               IP INTEL
    # ======================================================================
    IP_COLS = ["IP", "Public", "Events", "Users", "Failed sign-ins", "OK sign-ins", "First seen", "Last seen",
               "Country", "CC", "ASN", "ISP / Org", "Proxy/VPN", "Hosting", "Abuse score", "Abuse reports",
               "Tor", "Top users"]

    def _build_ip_table(self):
        agg = {}
        for rec in self.events:
            ip = clean_ip(ip_of(rec))
            if not ip:
                continue
            a = agg.setdefault(ip, {"n": 0, "users": Counter(), "fail": 0, "ok": 0, "first": "", "last": ""})
            a["n"] += 1
            if rec.get("UserId"):
                a["users"][str(rec["UserId"])] += 1
            op = _op(rec)
            a["fail"] += op in LOGIN_FAILED
            a["ok"] += op in LOGIN_OK
            ct = str(rec.get("CreationTime", ""))[:19]
            if ct:
                a["first"] = min(a["first"] or ct, ct)
                a["last"] = max(a["last"], ct)
        rows = []
        for ip, a in agg.items():
            public = is_public_ip(ip)
            if not public:
                self.ip_intel.setdefault(ip, {})["Country"] = "(non-public)"
            rows.append({"IP": ip, "ClientIP": ip, "Public": "yes" if public else "no", "Events": a["n"],
                         "Users": len(a["users"]), "Failed sign-ins": a["fail"], "OK sign-ins": a["ok"],
                         "First seen": a["first"], "Last seen": a["last"],
                         "Top users": ", ".join(u for u, _ in a["users"].most_common(3))})
        rows.sort(key=lambda r: (-r["Failed sign-ins"], -r["Events"]))
        self.ip_rows = rows
        self._render_ip_table()

    def _render_ip_table(self):
        for r in self.ip_rows:
            r.update(self.ip_intel.get(r["IP"], {}))
        rows = [[r.get(c) for c in self.IP_COLS] for r in self.ip_rows]
        public = sum(r["Public"] == "yes" for r in self.ip_rows)
        geo = sum(1 for r in self.ip_rows if r.get("CC"))
        label = (f"{len(self.ip_rows):,} IP addresses ({public:,} public) - {geo:,} with geo data - "
                 "sorted by failed sign-ins, then events")
        self.ip_grid.set_data(self.IP_COLS, rows, label, self.ip_rows)
        self._fill_summary()

    def geo_lookup(self):
        if not self.ip_rows:
            messagebox.showinfo("Geo lookup", "Fetch events first.")
            return
        ips = [r["IP"] for r in self.ip_rows if r["Public"] == "yes" and not self.ip_intel.get(r["IP"], {}).get("CC")]
        if not ips:
            messagebox.showinfo("Geo lookup", "All public IPs already have geo data.")
            return
        if not messagebox.askyesno(
                "Geo lookup", f"Send {len(ips):,} public IP address(es) to ip-api.com to look up country and "
                              "network owner?\n\nOnly IP addresses are sent - no user names or log data. The free "
                              "service uses plain HTTP and allows 15 requests of 100 IPs per minute.\n\nMake sure "
                              "this is allowed under your client agreement."):
            return
        fields = "status,message,query,country,countryCode,as,isp,org,proxy,hosting"

        def fn():
            found = {}
            for i in range(0, len(ips), 100):
                chunk = ips[i:i + 100]
                for attempt in range(3):
                    r = requests.post(GEO_LOOKUP_URL, params={"fields": fields}, json=chunk, timeout=30,
                                      verify=VERIFY_SSL, proxies=PROXIES)
                    if r.status_code != 429:
                        break
                    wait = int(r.headers.get("X-Ttl", 60)) + 1
                    self.log(f"ip-api.com rate limit - waiting {wait}s...")
                    self._wait(wait)
                if r.status_code != 200:
                    raise M365Error(f"Geo lookup failed: HTTP {r.status_code} {r.text[:200]}")
                for it in r.json():
                    if it.get("status") == "success":
                        found[it["query"]] = {"Country": it.get("country", ""), "CC": it.get("countryCode", ""),
                                              "ASN": it.get("as", ""), "ISP / Org": it.get("isp") or it.get("org", ""),
                                              "Proxy/VPN": "yes" if it.get("proxy") else "",
                                              "Hosting": "yes" if it.get("hosting") else ""}
                n = min(i + 100, len(ips))
                self.ui(lambda n=n: self.status.configure(text=f"Geo lookup {n}/{len(ips)}"))
                if n < len(ips) and r.headers.get("X-Rl") == "0":
                    self._wait(int(r.headers.get("X-Ttl", 60)) + 1)
            return found

        def done(found):
            for ip, info in found.items():
                self.ip_intel.setdefault(ip, {}).update(info)
            self._render_ip_table()
            self.activity.write("ip_geo_lookup", "ip-api.com", len(found), tenant=self._tenant())
            self.log(f"Geo lookup: {len(found)} of {len(ips)} IP(s) resolved")
            self.nb.select(self.tab_ip)
        self.run_bg(fn, done, f"Looking up {len(ips)} IP(s) on ip-api.com...")

    def forget_abuse_key(self):
        self._abuse_key = None
        self.log("AbuseIPDB key cleared from memory.")

    def abuse_check(self):
        if not self.ip_rows:
            messagebox.showinfo("AbuseIPDB", "Fetch events first.")
            return
        todo = [r["IP"] for r in self.ip_rows
                if r["Public"] == "yes" and "Abuse score" not in self.ip_intel.get(r["IP"], {})]
        if not todo:
            messagebox.showinfo("AbuseIPDB", "All public IPs have already been checked.")
            return
        if not self._abuse_key:
            key = simpledialog.askstring("AbuseIPDB", "AbuseIPDB API key (kept in memory only, never saved):",
                                         show="\u2022", parent=self.root)
            if not key or not key.strip():
                return
            self._abuse_key = key.strip()
        ips = todo[:ABUSEIPDB_MAX_IPS]
        if not messagebox.askyesno(
                "AbuseIPDB", f"Check {len(ips)} of {len(todo)} unchecked public IP(s) against AbuseIPDB?\n\n"
                             f"IPs with failed sign-ins go first. This uses {len(ips)} checks of your daily "
                             "AbuseIPDB quota (free plan: 1000/day)."):
            return
        key = self._abuse_key

        def fn():
            found, err = {}, ""
            for n, ip in enumerate(ips, 1):
                if self.stop.is_set():
                    err = "Cancelled"
                    break
                r = requests.get(ABUSEIPDB_URL, headers={"Key": key, "Accept": "application/json"},
                                 params={"ipAddress": ip, "maxAgeInDays": 90}, timeout=20,
                                 verify=VERIFY_SSL, proxies=PROXIES)
                if r.status_code in (401, 403):
                    err = "key"
                    break
                if r.status_code == 429:
                    err = "AbuseIPDB rate or daily limit reached - try again later."
                    break
                if r.status_code != 200:
                    err = f"AbuseIPDB returned HTTP {r.status_code} for {ip}"
                    break
                d = r.json().get("data", {})
                found[ip] = {"Abuse score": d.get("abuseConfidenceScore"), "Abuse reports": d.get("totalReports"),
                             "Tor": "yes" if d.get("isTor") else "", "_cc": d.get("countryCode") or "",
                             "_isp": d.get("isp") or ""}
                if n % 10 == 0:
                    self.ui(lambda n=n: self.status.configure(text=f"AbuseIPDB {n}/{len(ips)}"))
            return found, err

        def done(res):
            found, err = res
            for ip, info in found.items():
                cur = self.ip_intel.setdefault(ip, {})
                cc, isp = info.pop("_cc"), info.pop("_isp")
                cur.update(info)
                if not cur.get("CC") and cc:
                    cur["CC"] = cc
                if not cur.get("ISP / Org") and isp:
                    cur["ISP / Org"] = isp
            self._render_ip_table()
            self.activity.write("ip_abuseipdb", "AbuseIPDB", len(found), "ok" if not err else f"error: {err}",
                                tenant=self._tenant())
            self.log(f"AbuseIPDB: {len(found)} IP(s) checked")
            self.nb.select(self.tab_ip)
            if err == "key":
                self._abuse_key = None
                messagebox.showerror("AbuseIPDB", "The API key was rejected. It has been cleared - try again.")
            elif err and err != "Cancelled":
                messagebox.showwarning("AbuseIPDB", f"{err}\n\n{len(found)} IP(s) were checked before this.")
        self.run_bg(fn, done, f"Checking {len(ips)} IP(s) on AbuseIPDB...")

    def export_ips(self):
        recs = [{c: r.get(c) for c in self.IP_COLS} for r in self.ip_rows]
        self._export_records("m365_ip_intel", "IPIntel", self.IP_COLS, recs)

    # ======================================================================
    #                             USER TIMELINE
    # ======================================================================
    TL_COLS = ["CreationTime", "Workload", "Operation", "ResultStatus", "Client IP", "Country",
               "Object / detail", "Hunt flags", "MITRE", "UserAgent"]
    TL_SKIP_HUNTS = {"Successful sign-ins", "Failed operations (any workload)", "Mail items accessed",
                     "File downloads", "Mail deleted"}       # too noisy to flag every event

    def _fill_timeline_users(self):
        counts = Counter(str(e.get("UserId")) for e in self.events if e.get("UserId"))
        self.tl_cb.configure(values=[u for u, _ in counts.most_common(3000)])
        self.tl_grid.set_data([], [], "Pick a user and click 'Show timeline' (or right-click a user anywhere).")
        self.tl_rows = []

    @staticmethod
    def _detail(rec):
        for k in ("ObjectId", "SourceFileName", "SiteUrl", "MailboxOwnerUPN"):
            if rec.get(k):
                return str(rec[k])[:250]
        params = rec.get("Parameters")
        if isinstance(params, list):
            return "; ".join(f"{p.get('Name')}={p.get('Value')}" for p in params if isinstance(p, dict))[:250]
        return ""

    @staticmethod
    def _user_agent(rec):
        if rec.get("UserAgent"):
            return rec["UserAgent"]
        for p in rec.get("ExtendedProperties") or []:
            if isinstance(p, dict) and p.get("Name") == "UserAgent":
                return p.get("Value", "")
        return ""

    def show_timeline(self, user=None):
        if user is not None:
            self.tl_user.set(user)
        user = self.tl_user.get().strip()
        self.nb.select(self.tab_timeline)
        if not user:
            return
        if not self.events:
            messagebox.showinfo("User timeline", "Fetch events first.")
            return
        ul = user.lower()
        evs = [e for e in self.events if str(e.get("UserId", "")).lower() == ul]
        if not evs:
            evs = [e for e in self.events if ul in str(e.get("UserId", "")).lower()]
        evs.sort(key=lambda e: str(e.get("CreationTime", "")))
        hunts = [(h, EventFilter(operations=h.get("operations", ()), text=h.get("text", ()),
                                 result=h.get("result", "")))
                 for h in QUICK_HUNTS if h["name"] not in self.TL_SKIP_HUNTS]
        rows, recs = [], []
        for e in evs:
            matched = [h for h, f in hunts if f.match(e)]
            if self.tl_flagged.get() and not matched:
                continue
            ip = ip_of(e)
            row = {"CreationTime": e.get("CreationTime"), "Workload": e.get("Workload"),
                   "Operation": e.get("Operation"), "ResultStatus": e.get("ResultStatus"), "Client IP": ip,
                   "Country": self.ip_intel.get(clean_ip(ip), {}).get("Country", ""),
                   "Object / detail": self._detail(e), "Hunt flags": "; ".join(h["name"] for h in matched),
                   "MITRE": ", ".join(sorted({t for h in matched for t in h.get("mitre", [])})),
                   "UserAgent": self._user_agent(e)}
            rows.append(row)
            recs.append(e)
        self.tl_rows = [dict(r, **{"Full record": json.dumps(e, default=str)}) for r, e in zip(rows, recs)]
        flagged = sum(1 for r in rows if r["Hunt flags"])
        ips = {clean_ip(r["Client IP"]) for r in rows if r["Client IP"]}
        countries = {r["Country"] for r in rows if r["Country"] and r["Country"] != "(non-public)"}
        span = f"{rows[0]['CreationTime']} -> {rows[-1]['CreationTime']}" if rows else "no events"
        label = (f"{user}: {len(rows):,} event(s), {span}; {flagged} flagged by hunts; "
                 f"{len(ips)} IP(s); {len(countries)} countr{'y' if len(countries) == 1 else 'ies'}")
        self.tl_grid.set_data(self.TL_COLS, [[r[c] for c in self.TL_COLS] for r in rows], label, recs)
        self.activity.write("user_timeline", user, len(rows), tenant=self._tenant())

    def export_timeline(self):
        if not self.tl_rows:
            messagebox.showinfo("Export", "Show a user timeline first.")
            return
        user = self.tl_user.get().strip()
        self._export_records(f"m365_timeline_{safe_name(user)[:60]}", "Timeline", self.TL_COLS + ["Full record"],
                             self.tl_rows, {"user": user})

    # ======================================================================
    #                          SHARED EXPORT HELPER
    # ======================================================================
    def _export_records(self, prefix, sheet, cols, recs, extra=None):
        if not recs:
            messagebox.showinfo("Export", "Nothing to export.")
            return
        now = datetime.now(timezone.utc).replace(microsecond=0)
        start, end = self.fetch_window or (now, now)
        fmt, outdir, zip_it = self.fmt.get(), self._outdir(), self.zip_var.get()
        meta = self._meta(start, end, fmt, self.fetch_types)
        meta.update({"window_note": self.fetch_note or "none", "failed_blobs": self.fetch_failed,
                     "complete": self.fetch_failed == 0, "export": sheet})
        meta.update(extra or {})
        name = f"{prefix}_{stamp(start, end)}"
        recs = list(recs)

        def fn():
            batch = OutputBatch(outdir, name, fmt, zip_it, meta)
            sink = batch.sink(sheet)
            sink.write_records(recs, cols)
            batch.add(sheet, sink)
            files = batch.finish()
            self.activity.write(f"export_{sheet.lower()}", "; ".join(files), len(recs), tenant=self._tenant())
            return files

        def done(files):
            for f in files:
                self.log(f"Exported -> {f}")
            messagebox.showinfo("Export finished", f"{len(recs):,} rows\n\n" +
                                "\n".join(os.path.basename(f) for f in files) + f"\n\nFolder: {outdir}")
        self.run_bg(fn, done, f"Exporting {sheet}...")


def main():
    if sys.platform.startswith("win"):
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    root = tk.Tk()
    try:
        ttk.Style().theme_use("vista" if sys.platform.startswith("win") else "clam")
    except tk.TclError:
        pass
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
